from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from isite2.localization import excel_labels, scene_label
from isite2.output.ppt import scene_card_label
from isite2.rules.config_loader import scene_definitions


PPT_RECOMMENDATION_MARKERS = {
    "zh": "重点推荐物业点",
    "en": "Priority Recommended Sites",
}
RECOMMENDATION_SHEETS = {
    "zh": "推荐清单",
    "en": "Recommendations",
}
PROPERTY_HEADERS = {
    "zh": "物业名",
    "en": "Property Name",
}
SCENE_HEADERS = {
    "zh": "场景类型",
    "en": "Scene Type",
}
METRIC_VALUE_HEADERS = {
    "zh": "门槛指标值",
    "en": "Threshold Metric Value",
}
METRIC_KEY_HEADERS = {
    "zh": "门槛指标",
    "en": "Threshold Metric",
}
MAX_AIRPORT_ANNUAL_PASSENGERS = 250_000_000
MAX_HOTEL_KEYS = 5_000


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Validate that displayed one-page scene-card or legacy slide-2 "
            "recommendation metrics match the Excel recommendation sheet."
        )
    )
    parser.add_argument("--pptx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--locale", choices=["zh", "en"], required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    result = validate_ppt_metrics(args.pptx, args.xlsx, args.locale)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload, encoding="utf-8")
    print(payload)
    return 1 if result["failures"] else 0


def validate_ppt_metrics(pptx: Path, xlsx: Path, locale: str) -> dict:
    expected = _excel_recommendations(xlsx, locale)
    observed = _ppt_recommendation_rows(pptx, locale)
    failures = _excel_metric_sanity_failures(expected, locale)
    for row in observed:
        key = _entry_key(row["property_name"], row["scene"])
        expected_row = expected.get(key)
        if expected_row is None:
            failures.append({**row, "issue": "ppt_recommendation_missing_from_excel"})
            continue
        observed_value = _parse_number(row.get("metric_text"))
        expected_value = _parse_number(expected_row.get("metric_value"))
        if observed_value != expected_value:
            failures.append(
                {
                    **row,
                    "issue": "ppt_metric_value_mismatch",
                    "observed_value": observed_value,
                    "expected_value": expected_value,
                    "excel_metric_value": expected_row.get("metric_value"),
                }
            )
    summary = _ppt_summary(pptx, locale)
    if (
        summary["recommendation_total"] is not None
        and summary["recommendation_total"] != len(expected)
    ):
        failures.append(
            {
                "issue": "ppt_recommendation_total_mismatch",
                "observed_value": summary["recommendation_total"],
                "expected_value": len(expected),
            }
        )
    if (
        summary["card_recommendation_total"] is not None
        and summary["card_recommendation_total"]
        != _displayed_scene_recommendation_total(
            expected,
            summary["card_scenes"],
            locale,
        )
    ):
        expected_card_total = _displayed_scene_recommendation_total(
            expected,
            summary["card_scenes"],
            locale,
        )
        failures.append(
            {
                "issue": "ppt_scene_card_recommendation_total_mismatch",
                "observed_value": summary["card_recommendation_total"],
                "expected_value": expected_card_total,
            }
        )
    missing_in_ppt = [
        row
        for key, row in expected.items()
        if key not in {_entry_key(item["property_name"], item["scene"]) for item in observed}
    ]
    return {
        "pptx": str(pptx),
        "xlsx": str(xlsx),
        "locale": locale,
        "ppt_recommendation_count": len(observed),
        "excel_recommendation_count": len(expected),
        "missing_in_ppt_count": len(missing_in_ppt),
        "ppt_slide_count": summary["slide_count"],
        "ppt_recommendation_total": summary["recommendation_total"],
        "ppt_scene_card_recommendation_total": summary["card_recommendation_total"],
        "failures": failures,
    }


def validate_excel_metrics(
    xlsx: Path,
    locale: str,
    *,
    expected_country: str | None = None,
) -> dict:
    """Validate a standalone country workbook without requiring PPT assets."""
    labels = excel_labels(locale)
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    expected_sheets = [labels["sheets"][key] for key in labels["required_sheet_keys"]]
    missing_sheets = [name for name in expected_sheets if name not in workbook.sheetnames]
    failures: list[dict] = [
        {"issue": "excel_required_sheet_missing", "sheet": name}
        for name in missing_sheets
    ]

    recommendations: dict[tuple[str, str], dict] = {}
    recommendation_sheet = RECOMMENDATION_SHEETS[locale]
    if recommendation_sheet in workbook.sheetnames:
        recommendations = _excel_recommendations(xlsx, locale)
        failures.extend(_excel_metric_sanity_failures(recommendations, locale))

    countries: set[str] = set()
    row_count = 0
    main_sheet = labels["sheets"]["main"]
    if main_sheet in workbook.sheetnames:
        worksheet = workbook[main_sheet]
        headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        country_col = _header_index(headers, labels["columns"]["main"][0])
        property_col = _header_index(headers, labels["columns"]["main"][2])
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            property_name = str(row[property_col] or "").strip()
            if not property_name:
                continue
            row_count += 1
            country = str(row[country_col] or "").strip()
            if country:
                countries.add(country)
        if row_count == 0:
            failures.append({"issue": "excel_main_table_empty"})
        if len(countries) > 1:
            failures.append(
                {
                    "issue": "excel_country_scope_mixed",
                    "countries": sorted(countries),
                }
            )
        if expected_country and {item.casefold() for item in countries} != {
            expected_country.casefold()
        }:
            failures.append(
                {
                    "issue": "excel_country_scope_mismatch",
                    "observed_countries": sorted(countries),
                    "expected_country": expected_country,
                }
            )
    workbook.close()
    return {
        "xlsx": str(xlsx),
        "locale": locale,
        "main_table_row_count": row_count,
        "recommendation_count": len(recommendations),
        "countries": sorted(countries),
        "missing_required_sheets": missing_sheets,
        "failures": failures,
    }


def _excel_metric_sanity_failures(
    expected: dict[tuple[str, str], dict],
    locale: str,
) -> list[dict]:
    airport_labels = {
        scene_label("airport_terminal", locale).casefold(),
        scene_card_label("airport_terminal", locale).casefold(),
    }
    hotel_labels = {
        scene_label("luxury_hotel_mice", locale).casefold(),
        scene_card_label("luxury_hotel_mice", locale).casefold(),
    }
    failures: list[dict] = []
    for row in expected.values():
        scene = str(row.get("scene") or "").strip().casefold()
        metric_key = str(row.get("metric_key") or "").strip().casefold()
        metric_value = _parse_number(row.get("metric_value"))
        if metric_value is None:
            continue
        if (
            scene in airport_labels
            and metric_key
            in {"annual_passenger_throughput", "passenger_throughput"}
            and metric_value > MAX_AIRPORT_ANNUAL_PASSENGERS
        ):
            failures.append(
                {
                    "property_name": row["property_name"],
                    "scene": row["scene"],
                    "issue": "excel_airport_metric_implausible_high",
                    "observed_value": metric_value,
                    "maximum_plausible_value": MAX_AIRPORT_ANNUAL_PASSENGERS,
                    "excel_metric_value": row.get("metric_value"),
                }
            )
        if (
            scene in hotel_labels
            and metric_key in {"keys", "room_count"}
            and metric_value > MAX_HOTEL_KEYS
        ):
            failures.append(
                {
                    "property_name": row["property_name"],
                    "scene": row["scene"],
                    "issue": "excel_hotel_keys_implausible_high",
                    "observed_value": metric_value,
                    "maximum_plausible_value": MAX_HOTEL_KEYS,
                    "excel_metric_value": row.get("metric_value"),
                }
            )
    return failures


def _excel_recommendations(xlsx: Path, locale: str) -> dict[tuple[str, str], dict]:
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    sheet_name = RECOMMENDATION_SHEETS[locale]
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"{xlsx} missing sheet {sheet_name}")
    worksheet = workbook[sheet_name]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    property_col = _header_index(headers, PROPERTY_HEADERS[locale])
    scene_col = _header_index(headers, SCENE_HEADERS[locale])
    metric_key_col = _header_index(headers, METRIC_KEY_HEADERS[locale])
    metric_value_col = _header_index(headers, METRIC_VALUE_HEADERS[locale])
    rows: dict[tuple[str, str], dict] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        property_name = str(row[property_col] or "").strip()
        scene = str(row[scene_col] or "").strip()
        if not property_name:
            continue
        rows[_entry_key(property_name, scene)] = {
            "property_name": property_name,
            "scene": scene,
            "metric_key": row[metric_key_col],
            "metric_value": row[metric_value_col],
        }
    return rows


def _ppt_recommendation_rows(pptx: Path, locale: str) -> list[dict]:
    deck = Presentation(pptx)
    if len(deck.slides) == 1:
        return _one_page_recommendation_rows(deck, locale)
    texts = _slide_texts(pptx, "ppt/slides/slide2.xml")
    marker = PPT_RECOMMENDATION_MARKERS[locale]
    try:
        start = texts.index(marker) + 1
    except ValueError as exc:
        raise ValueError(f"{pptx} slide 2 missing marker {marker!r}") from exc

    rows: list[dict] = []
    index = start
    while index + 2 < len(texts):
        name_text, scene, metric_text = texts[index : index + 3]
        match = re.match(r"^\s*(\d+)\.\s+(.+?)\s*$", name_text)
        if not match:
            break
        rows.append(
            {
                "rank": int(match.group(1)),
                "property_name": match.group(2),
                "scene": scene,
                "metric_text": metric_text,
            }
        )
        index += 3
    return rows


def _one_page_recommendation_rows(
    deck: Presentation,
    locale: str,
) -> list[dict]:
    shapes = {
        shape.name: shape.text.strip()
        for shape in deck.slides[0].shapes
        if shape.name and hasattr(shape, "text")
    }
    qualified = {"达标"} if locale == "zh" else {"QUALIFIED", "Qualified"}
    rows: list[dict] = []
    for index in range(1, 10):
        suffix = f"{index:02d}"
        property_name = shapes.get(f"name-{suffix}", "").strip()
        status = shapes.get(f"status-{suffix}", "").strip()
        if not property_name or status not in qualified:
            continue
        rows.append(
            {
                "rank": index,
                "property_name": property_name,
                "scene": shapes.get(f"scene-{suffix}", "").strip(),
                "metric_text": shapes.get(f"metric-{suffix}", "").strip(),
            }
        )
    return rows


def _ppt_summary(pptx: Path, locale: str) -> dict:
    deck = Presentation(pptx)
    summary = {
        "slide_count": len(deck.slides),
        "recommendation_total": None,
        "card_recommendation_total": None,
        "card_scenes": [],
    }
    if len(deck.slides) != 1:
        return summary
    shapes = {
        shape.name: shape.text.strip()
        for shape in deck.slides[0].shapes
        if shape.name and hasattr(shape, "text")
    }
    subtitle = shapes.get("subtitle", "")
    pattern = (
        r"(\d+)\s*个门槛推荐"
        if locale == "zh"
        else r"(\d+)\s*threshold recommendations"
    )
    match = re.search(pattern, subtitle, flags=re.IGNORECASE)
    if match:
        summary["recommendation_total"] = int(match.group(1))
    card_total = 0
    card_count = 0
    for index in range(1, 10):
        scene = shapes.get(f"scene-{index:02d}", "").strip()
        if scene:
            summary["card_scenes"].append(scene)
        counts = shapes.get(f"counts-{index:02d}", "")
        if not counts:
            continue
        card_match = re.search(
            r"(?:推荐|Rec\.)\s*(\d+)",
            counts,
            flags=re.IGNORECASE,
        )
        if card_match:
            card_count += 1
            card_total += int(card_match.group(1))
    if card_count:
        summary["card_recommendation_total"] = card_total
    return summary


def _displayed_scene_recommendation_total(
    expected: dict[tuple[str, str], dict],
    card_scenes: list[str],
    locale: str,
) -> int:
    excel_scene_labels: set[str] = set()
    for card_scene in card_scenes:
        matched = False
        for scene_type in scene_definitions():
            if scene_card_label(scene_type, locale) != card_scene:
                continue
            excel_scene_labels.add(scene_label(scene_type, locale))
            matched = True
            break
        if not matched:
            excel_scene_labels.add(card_scene)
    return sum(
        1
        for row in expected.values()
        if str(row.get("scene") or "").strip() in excel_scene_labels
    )


def _slide_texts(pptx: Path, slide_name: str) -> list[str]:
    with zipfile.ZipFile(pptx) as archive:
        root = ET.fromstring(archive.read(slide_name))
    return [
        node.text or ""
        for node in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t")
    ]


def _header_index(headers: list[str], header: str) -> int:
    try:
        return headers.index(header)
    except ValueError as exc:
        raise ValueError(f"missing header {header!r}; headers={headers}") from exc


def _entry_key(property_name: str, scene: str) -> tuple[str, str]:
    return (property_name.casefold().strip(), _normalize_scene_key(scene))


def _normalize_scene_key(scene: str) -> str:
    normalized = re.sub(r"[\s_/\-]+", " ", str(scene or "").casefold()).strip()
    aliases = (
        (("airport", "机场"), "airport_terminal"),
        (("transport", "transit", "railway", "交通枢纽"), "transport_hub"),
        (("convention", "exhibition", "会展"), "convention_center"),
        (("mall", "retail", "商场", "商超"), "mall_mixed_use"),
        (("stadium", "体育场"), "stadium"),
        (("hotel", "酒店"), "luxury_hotel_mice"),
        (("office", "写字楼", "办公"), "office_government"),
        (("hospital", "医院"), "hospital"),
        (("university", "大学"), "university"),
        (("mosque", "清真寺"), "mosque"),
    )
    for tokens, canonical in aliases:
        if any(token in normalized for token in tokens):
            return canonical
    return normalized


def _parse_number(value: object) -> int | None:
    if value is None:
        return None
    text = str(value)
    match = re.search(r"\d[\d,]*(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group(0).replace(",", ""))
    return int(number) if number.is_integer() else round(number)


if __name__ == "__main__":
    sys.exit(main())
