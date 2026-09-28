import json

from openpyxl import load_workbook

from isite2.growth.africa_loop import (
    AFRICAN_COUNTRIES,
    registry_backed_african_countries,
    run_africa_scan_round,
)
from isite2.repositories.sqlalchemy import SQLAlchemyScanRunRepository
from isite2.rules.config_loader import load_source_registry


def test_registry_backed_african_countries_only_returns_seeded_countries() -> None:
    countries = registry_backed_african_countries(load_source_registry())

    assert countries == ["Algeria", "Egypt"]
    assert set(countries).issubset(set(AFRICAN_COUNTRIES))


def test_africa_loop_round_persists_database_artifacts_and_report(tmp_path) -> None:
    repository = SQLAlchemyScanRunRepository.from_url(
        f"sqlite+pysqlite:///{tmp_path / 'isite2_loop.db'}",
        storage_mode="sqlite",
    )

    first_round = run_africa_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
    )

    assert first_round.round_number == 1
    assert first_round.country == "Algeria"
    assert first_round.candidate_count == 4
    assert first_round.review_count == 4
    assert first_round.report_path.exists()
    assert first_round.summary_path.exists()
    assert first_round.excel_path.exists()
    assert (tmp_path / "loop" / "latest_report.md").exists()

    workbook = load_workbook(first_round.excel_path)
    assert workbook["Main"].max_row == 5

    summary = json.loads(first_round.summary_path.read_text(encoding="utf-8"))
    assert summary["country"] == "Algeria"
    assert summary["ui_paths"]["geojson"].endswith("country=Algeria")
    assert "Source registry coverage is 2/54" in summary["improvement_points"][0]
    assert len(repository.list_properties({"country": "Algeria"})) == 4

    artifacts = repository.list_output_artifacts(first_round.run_id)
    assert {artifact["artifact_type"] for artifact in artifacts} == {
        "excel",
        "json_summary",
        "markdown_report",
    }

    second_round = run_africa_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
    )

    state = json.loads((tmp_path / "state.json").read_text(encoding="utf-8"))
    assert second_round.round_number == 2
    assert second_round.country == "Egypt"
    assert state["registry_country_index"] == 0
    assert state["cycles_completed"] == 1
    assert len(state["history"]) == 2
    assert len(repository.list_properties({"country": "Egypt"})) == 4
