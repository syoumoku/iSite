import json

from openpyxl import load_workbook

from isite2.growth.evidence_intake import (
    CandidateDraft,
    EvidenceIntakeProvider,
    SeedCatalogEvidenceIntakeProvider,
    validate_candidate_draft,
)
from isite2.growth.regional_loop import (
    ASIA_PACIFIC_COUNTRIES,
    DEFAULT_REGIONS,
    LATIN_AMERICA_COUNTRIES,
    NORTH_AFRICA_COUNTRIES,
    countries_for_regions,
    registry_backed_country_counts,
    run_regional_scan_round,
    select_country_batch,
)
from isite2.orchestrator.pipeline import _numeric_metric_value
from isite2.repositories.sqlalchemy import SQLAlchemyScanRunRepository
from isite2.rules.candidate_quality import filter_packets_for_surface
from isite2.rules.config_loader import load_source_registry


class EmptyIntakeProvider(EvidenceIntakeProvider):
    def discover(self, regions: list[str], target_countries: list[str]) -> list[CandidateDraft]:
        return []


def test_default_regions_cover_africa_and_latin_america_country_lists() -> None:
    countries = countries_for_regions(DEFAULT_REGIONS)

    assert "Algeria" in countries
    assert "Brazil" in countries
    assert "Mexico" in countries
    assert len(LATIN_AMERICA_COUNTRIES) == 33
    assert len(countries) == 87


def test_registry_metric_parser_ignores_reporting_year_before_million_value() -> None:
    assert (
        _numeric_metric_value("January-May 2025 passenger movement: 4.9 million passengers.")
        == 4_900_000
    )
    assert (
        _numeric_metric_value("2025 annual passenger traffic: 47.2 million passengers.")
        == 47_200_000
    )
    assert (
        _numeric_metric_value("annual_passenger_throughput: 43.712 million passengers in 2024")
        == 43_712_000
    )


def test_asia_pacific_region_includes_initial_scan_batch_countries() -> None:
    countries = countries_for_regions(["Asia Pacific"])

    assert ASIA_PACIFIC_COUNTRIES == [
        "Sri Lanka",
        "Cambodia",
        "Maldives",
        "Philippines",
        "Vietnam",
        "Indonesia",
        "Thailand",
    ]
    assert countries == [
        "Sri Lanka",
        "Cambodia",
        "Maldives",
        "Philippines",
        "Vietnam",
        "Indonesia",
        "Thailand",
    ]


def test_north_africa_region_matches_business_scan_scope() -> None:
    countries = countries_for_regions(["North Africa"])

    assert NORTH_AFRICA_COUNTRIES == [
        "Egypt",
        "Ethiopia",
        "Algeria",
        "Morocco",
        "Cameroon",
        "Senegal",
        "Cote d'Ivoire",
        "Congo",
        "Mali",
        "Burkina Faso",
        "Guinea",
        "Gambia",
        "Mauritania",
        "Libya",
        "Tunisia",
        "Democratic Republic of the Congo",
        "Gabon",
        "Chad",
        "Equatorial Guinea",
        "Central African Republic",
        "Cape Verde",
        "Benin",
    ]
    assert countries == NORTH_AFRICA_COUNTRIES
    assert "North Africa" not in DEFAULT_REGIONS


def test_registry_backed_counts_include_only_objective_seeded_countries() -> None:
    counts = registry_backed_country_counts(DEFAULT_REGIONS, load_source_registry())

    assert [item["country"] for item in counts] == ["Algeria", "Egypt"]
    assert [item["candidate_count"] for item in counts] == [4, 4]


def test_select_country_batch_targets_30_to_50_when_registry_allows() -> None:
    registry_counts = [
        {"region": "Africa", "country": "A", "candidate_count": 12},
        {"region": "Africa", "country": "B", "candidate_count": 14},
        {"region": "Latin America", "country": "C", "candidate_count": 16},
        {"region": "Latin America", "country": "D", "candidate_count": 20},
    ]

    batch = select_country_batch({}, registry_counts, batch_min=30, batch_max=50)

    assert [item["country"] for item in batch] == ["A", "B", "C"]
    assert sum(item["candidate_count"] for item in batch) == 42


def test_regional_loop_round_persists_batch_report_and_artifacts(tmp_path) -> None:
    repository = SQLAlchemyScanRunRepository.from_url(
        f"sqlite+pysqlite:///{tmp_path / 'isite2_loop.db'}",
        storage_mode="sqlite",
    )

    first_round = run_regional_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
        regions=DEFAULT_REGIONS,
        batch_min=30,
        batch_max=50,
        overlay_path=tmp_path / "overlay.yaml",
        draft_path=tmp_path / "drafts.json",
        intake_provider=SeedCatalogEvidenceIntakeProvider(),
    )

    assert first_round.round_number == 1
    assert first_round.mode == "scan"
    assert first_round.scan_created is True
    assert first_round.candidate_count == 32
    assert first_round.review_count >= 32
    assert first_round.report_path.exists()
    assert first_round.summary_path.exists()
    assert first_round.excel_path.exists()
    assert (tmp_path / "loop" / "latest_report.md").exists()

    workbook = load_workbook(first_round.excel_path)
    surface_packets = filter_packets_for_surface(
        repository.list_properties({"scan_run_id": first_round.run_id}),
        "main_table",
    )
    assert workbook["Main"].max_row == len(surface_packets) + 1

    summary = json.loads(first_round.summary_path.read_text(encoding="utf-8"))
    assert summary["regions"] == DEFAULT_REGIONS
    assert summary["countries"][:2] == ["Algeria", "Egypt"]
    assert summary["batch_target"] == {"min": 30, "max": 50}
    assert summary["target_country_count"] == 87
    assert summary["registry_backed_country_count"] == 10
    assert summary["registry_backed_candidate_count"] == 40
    assert summary["intake"]["accepted_count"] == 32
    assert len(repository.list_properties({"country": "Algeria"})) == 4
    assert len(repository.list_properties({"country": "Egypt"})) == 4
    assert len(repository.list_properties({"country": "Nigeria"})) == 4

    artifacts = repository.list_output_artifacts(first_round.run_id)
    assert {artifact["artifact_type"] for artifact in artifacts} == {
        "excel",
        "json_summary",
        "markdown_report",
    }

    state = json.loads((tmp_path / "state.json").read_text(encoding="utf-8"))
    assert state["registry_item_index"] == 8
    assert state["cycles_completed"] == 0
    assert len(state["history"]) == 1


def test_regional_loop_enters_intake_only_mode_when_evidence_pool_underfilled(tmp_path) -> None:
    repository = SQLAlchemyScanRunRepository.from_url(
        f"sqlite+pysqlite:///{tmp_path / 'isite2_loop.db'}",
        storage_mode="sqlite",
    )

    round_result = run_regional_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
        regions=DEFAULT_REGIONS,
        batch_min=30,
        batch_max=50,
        overlay_path=tmp_path / "overlay.yaml",
        draft_path=tmp_path / "drafts.json",
        intake_provider=EmptyIntakeProvider(),
    )

    assert round_result.mode == "intake_only"
    assert round_result.scan_created is False
    assert round_result.run_id is None
    assert round_result.excel_path is None
    assert repository.list() == []
    summary = json.loads(round_result.summary_path.read_text(encoding="utf-8"))
    assert summary["status"] == "intake_underfilled"
    assert summary["registry_backed_candidate_count"] == 8


def test_regional_loop_idles_without_scan_when_no_new_evidence(tmp_path) -> None:
    repository = SQLAlchemyScanRunRepository.from_url(
        f"sqlite+pysqlite:///{tmp_path / 'isite2_loop.db'}",
        storage_mode="sqlite",
    )

    first_round = run_regional_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
        regions=DEFAULT_REGIONS,
        batch_min=30,
        batch_max=50,
        overlay_path=tmp_path / "overlay.yaml",
        draft_path=tmp_path / "drafts.json",
        intake_provider=SeedCatalogEvidenceIntakeProvider(),
    )
    second_round = run_regional_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
        regions=DEFAULT_REGIONS,
        batch_min=30,
        batch_max=50,
        overlay_path=tmp_path / "overlay.yaml",
        draft_path=tmp_path / "drafts.json",
        intake_provider=SeedCatalogEvidenceIntakeProvider(),
    )

    assert first_round.scan_created is True
    assert second_round.mode == "idle_no_new_evidence"
    assert second_round.scan_created is False
    assert len(repository.list()) == 1
    summary = json.loads(second_round.summary_path.read_text(encoding="utf-8"))
    assert summary["status"] == "no_new_evidence"
    assert summary["intake"]["new_evidence_count"] == 0
    assert summary["intake"]["duplicate_unchanged_count"] == 32


def test_regional_loop_targets_egypt_and_force_scans_existing_pool(tmp_path) -> None:
    repository = SQLAlchemyScanRunRepository.from_url(
        f"sqlite+pysqlite:///{tmp_path / 'isite2_loop.db'}",
        storage_mode="sqlite",
    )

    round_result = run_regional_scan_round(
        repository=repository,
        state_path=tmp_path / "state.json",
        output_dir=tmp_path / "loop",
        regions=["Africa"],
        batch_min=1,
        batch_max=50,
        overlay_path=tmp_path / "overlay.yaml",
        draft_path=tmp_path / "drafts.json",
        intake_provider=EmptyIntakeProvider(),
        target_countries=["Egypt"],
        max_searches_per_cycle=20,
        max_fetches_per_cycle=40,
        force_scan_existing_pool=True,
    )

    summary = json.loads(round_result.summary_path.read_text(encoding="utf-8"))

    assert round_result.mode == "scan"
    assert round_result.scan_created is True
    assert round_result.countries == ["Egypt"]
    assert round_result.candidate_count == 4
    assert summary["countries"] == ["Egypt"]
    assert summary["target_country_count"] == 1
    assert summary["registry_backed_country_count"] == 1
    assert summary["intake"]["searched_count"] <= 20
    assert summary["intake"]["fetched_count"] <= 40
    assert len(repository.list_properties({"country": "Egypt"})) == 4
    assert repository.list_properties({"country": "Algeria"}) == []


def test_candidate_draft_with_missing_evidence_is_rejected() -> None:
    draft = CandidateDraft(
        region="Africa",
        country="Nigeria",
        city="Lagos",
        property_name="Incomplete Draft",
        scene_type="airport_terminal",
        annual_visits=1,
        latitude=6.0,
        longitude=3.0,
        geocode_precision="venue centroid",
        map_source="public map",
        map_source_date="2026-05-08",
        field_group="airport_role",
        indicator_name="gateway_role",
        field_value="",
        source_name="",
        source_tier="Tier 3",
        source_url="",
        source_date="",
        bbox={
            "min_latitude": 4.0,
            "max_latitude": 14.2,
            "min_longitude": 2.5,
            "max_longitude": 15.0,
        },
    )

    validation = validate_candidate_draft(draft)

    assert validation.accepted is False
    assert "field_value missing" in validation.issues
    assert "source_url missing" in validation.issues


def test_candidate_draft_allows_missing_annual_visits_for_proxy_model() -> None:
    draft = CandidateDraft(
        region="Africa",
        country="Ghana",
        city="Accra",
        property_name="Accra Sports Stadium",
        scene_type="stadium",
        annual_visits=None,
        latitude=5.5529,
        longitude=-0.1914,
        geocode_precision="stadium venue centroid",
        map_source="public map",
        map_source_date="2026-05-15",
        field_group="seat_count",
        indicator_name="seat_count",
        field_value="40,000 seats",
        source_name="Example Stadium Authority",
        source_tier="Tier 1",
        source_url="https://example.org/accra-stadium-capacity",
        source_date="2026",
        bbox={
            "min_latitude": 4.5,
            "max_latitude": 11.5,
            "min_longitude": -3.5,
            "max_longitude": 1.5,
        },
    )

    validation = validate_candidate_draft(draft)

    assert validation.accepted is True
    assert validation.issues == []


def test_candidate_draft_city_page_pollution_is_rejected() -> None:
    draft = CandidateDraft(
        region="Africa",
        country="Malawi",
        city="Zomba",
        property_name="Beijing",
        scene_type="transport_hub",
        annual_visits=25,
        latitude=-15.3860585,
        longitude=35.3191201,
        geocode_precision="gift",
        map_source="OpenStreetMap Nominatim",
        map_source_date="2026-05-08",
        field_group="line_count",
        indicator_name="line_count",
        field_value="25 lines",
        source_name="en.wikipedia.org",
        source_tier="Tier 3",
        source_url="https://en.wikipedia.org/wiki/Beijing",
        source_date="2026-05-08",
        bbox={
            "min_latitude": -17.2,
            "max_latitude": -9.3,
            "min_longitude": 32.6,
            "max_longitude": 35.9,
        },
    )

    validation = validate_candidate_draft(draft)

    assert validation.accepted is False
    assert any("candidate_quality" in issue for issue in validation.issues)
    assert any("location-only" in issue for issue in validation.issues)
