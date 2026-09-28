from uuid import uuid4
from types import SimpleNamespace

from pptx import Presentation

from isite2.domain.enums import (
    ActionClass,
    BuildEvidenceStatus,
    CrossCheckStatus,
    EvidenceStatus,
    EvidenceType,
    IndoorRAT,
    IndoorSystemPresence,
    IndoorSystemType,
    ProxyLevel,
    RecommendedSolution,
    SceneForm,
    SourceTier,
    ValueClass,
)
from isite2.domain.models import (
    BuildStatus,
    Conclusion,
    EvidenceItem,
    PropertyEntity,
    PropertyHeroImage,
    SceneModelResult,
    SitePacket,
)
from isite2.output.excel import _build_recommendation_index
from isite2.output.ppt import (
    _card_metric_font_size,
    _image_format_from_magic,
    scene_card_scene_types,
    scene_card_packets,
    write_ppt_deck,
)
from isite2.rules.config_loader import scene_definitions


def test_ppt_standard_report_uses_one_page_template(tmp_path) -> None:
    output = tmp_path / "standard_report.pptx"
    packets = [
        _packet(
            scene_type="airport_terminal",
            property_name="Airport Alpha",
            strong_metric="annual_passenger_throughput",
            strong_value="2025 annual passenger traffic: 4,200,000 passengers.",
        ),
        _packet(
            scene_type="stadium",
            property_name="Stadium Bravo",
            strong_metric="seat_count",
            strong_value="Stadium capacity: 42,000 seats.",
        ),
        _packet(
            scene_type="airport_terminal",
            property_name="Airport Capacity Delta",
            strong_metric="terminal_capacity",
            strong_value=(
                "terminal_capacity: 3,000,000 passengers per year planned handling "
                "capacity; new passenger terminal of 32,000 m2"
            ),
        ),
        _packet(
            scene_type="luxury_hotel_mice",
            property_name="Hotel Echo",
            strong_metric="keys",
            strong_value="keys: 187 guest rooms; 1,000 m² conference center",
        ),
    ]

    write_ppt_deck(output, packets=packets, locale="en")

    deck = Presentation(output)
    assert len(deck.slides) == 1
    text = "\n".join(
        shape.text
        for slide in deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert "Directory leads" in text
    assert "Location & primary metric" in text
    assert "Proxy value model" in text
    assert "Quality gate" in text
    assert "4 qualified candidates" in text
    assert "Airport Alpha" in text
    assert "Stadium Bravo" in text
    assert "Hotel Echo" in text
    assert "Airport Capacity Delta" not in text
    assert "Annual passenger throughput: 4,200,000" in text
    assert "Terminal capacity: 32,000,000,000" not in text
    assert "Rooms / keys: 187" in text
    assert "Rooms / keys: 1,000,000,000" not in text
    assert "Rule / Advancement Logic" not in text
    assert "Location:" not in text


def test_ppt_image_magic_distinguishes_mpo_and_webp() -> None:
    assert _image_format_from_magic(b"\xff\xd8\xff\xe2MPF\x00payload") == "MPO"
    assert (
        _image_format_from_magic(b"RIFF\x10\x00\x00\x00WEBPVP8 payload")
        == "WEBP"
    )


def test_dense_scene_card_shrinks_long_english_metric_to_one_line() -> None:
    assert _card_metric_font_size(
        "Annual passenger throughput: 199,453",
        locale="en",
        dense=True,
    ) == 9.25
    assert _card_metric_font_size(
        "Seat count: 35,000",
        locale="en",
        dense=True,
    ) == 12.5


def test_one_page_scene_selection_keeps_full_pool_but_prioritizes_nine_cards() -> None:
    scene_types = [
        "cruise_port",
        "mosque",
        "university",
        "hospital",
        "office_government",
        "convention_center",
        "luxury_hotel_mice",
        "stadium",
        "mall_mixed_use",
        "transport_hub",
        "airport_terminal",
    ]
    packets = [
        SimpleNamespace(entity=SimpleNamespace(scene_type=scene_type))
        for scene_type in scene_types
    ]

    selected = scene_card_scene_types(packets)

    assert selected == [
        "airport_terminal",
        "transport_hub",
        "mall_mixed_use",
        "stadium",
        "luxury_hotel_mice",
        "convention_center",
        "office_government",
        "hospital",
        "university",
    ]
    assert len(packets) == 11


def test_scene_card_uses_ranked_image_fallback(tmp_path) -> None:
    preferred = _packet(
        scene_type="stadium",
        property_name="Preferred Stadium",
        strong_metric="seat_count",
        strong_value="Stadium capacity: 42,000 seats.",
        hero_url="https://example.com/preferred.jpg",
    )
    fallback = _packet(
        scene_type="stadium",
        property_name="Fallback Stadium",
        strong_metric="seat_count",
        strong_value="Stadium capacity: 30,000 seats.",
        hero_url="https://example.com/fallback.jpg",
    )
    recommendations = _build_recommendation_index(
        [preferred, fallback],
        locale="en",
    )

    selected = scene_card_packets(
        [preferred, fallback],
        recommendations,
        image_assets={"https://example.com/fallback.jpg": tmp_path / "fallback.jpg"},
        require_images=True,
    )

    assert selected[0].entity.property_name == "Fallback Stadium"


def _packet(
    *,
    scene_type: str,
    property_name: str,
    strong_metric: str,
    strong_value: str,
    hero_url: str | None = None,
) -> SitePacket:
    property_id = uuid4()
    scene_rule = scene_definitions()[scene_type]
    scene_form = (
        SceneForm.SEMI_OPEN
        if scene_rule["scene_form"] == SceneForm.SEMI_OPEN.value
        else SceneForm.INDOOR
    )
    return SitePacket(
        entity=PropertyEntity(
            property_id=property_id,
            country="Brazil",
            city="São Paulo",
            property_name=property_name,
            scene_type=scene_type,
            scene_form=scene_form,
            latitude=-23.5,
            longitude=-46.6,
            geocode_precision="Property",
            google_maps_link="https://www.google.com/maps/search/?api=1&query=-23.5,-46.6",
            hero_image=(
                PropertyHeroImage(
                    url=hero_url,
                    alt_text=property_name,
                    source_name="Official property website",
                    source_url=hero_url,
                )
                if hero_url
                else None
            ),
        ),
        scene=SceneModelResult(
            area_metric_name=scene_rule["area_metric"],
            primary_value_indicators=list(scene_rule["primary_indicators"]),
            proxy_basis=scene_rule["proxy_basis"][0],
            proxy_level=ProxyLevel.P1_STRONG,
            annual_visits_est=4_200_000,
        ),
        evidence=[
            EvidenceItem(
                property_id=property_id,
                field_group=strong_metric,
                indicator_name=strong_metric,
                field_value=strong_value,
                source_name="Official statistics",
                source_tier=SourceTier.TIER_1,
                source_url="https://example.com/statistics",
                evidence_type=EvidenceType.DIRECT,
                cross_check_status=CrossCheckStatus.PARTIAL,
            )
        ],
        build_status=BuildStatus(
            indoor_system_presence=IndoorSystemPresence.NO_PUBLIC_EVIDENCE,
            indoor_system_type=IndoorSystemType.UNKNOWN,
            indoor_rat=IndoorRAT.UNKNOWN,
            build_evidence_status=BuildEvidenceStatus.UNKNOWN,
        ),
        conclusion=Conclusion(
            evidence_status=EvidenceStatus.SUPPORTED,
            value_class=ValueClass.CITY_CORE,
            action_class=ActionClass.SURVEY_FIRST,
            recommended_solution=RecommendedSolution.PRRU,
            reason_to_recommend="Hard primary metric supports the site recommendation.",
            next_action="Verify indoor coverage status and operator evidence.",
        ),
    )
