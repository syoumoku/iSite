from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from scripts import generate_standard_country_report as report_generation
from scripts import qa_standard_report_output_bundle as report_audit
from scripts.qa_standard_report_output_bundle import audit_report_dir
from isite2.localization import excel_labels


def test_standard_report_output_audit_passes_matching_ppt_and_excel(tmp_path: Path) -> None:
    _write_excel(tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx")
    _write_ppt(
        tmp_path / "isite_test_standard_ppt_en_20260707T000000.pptx",
        airport_metric="Terminal capacity: 4,200,000",
        hotel_metric="Rooms / keys: 187",
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
    )

    assert result["status"] == "pass"
    assert result["rule_audits"]["en"]["failures"] == []


def test_standard_report_output_audit_blocks_metric_mismatch(tmp_path: Path) -> None:
    _write_excel(tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx")
    _write_ppt(
        tmp_path / "isite_test_standard_ppt_en_20260707T000000.pptx",
        airport_metric="Terminal capacity: 4,200,000",
        hotel_metric="Rooms / keys: 1,000,000",
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
    )

    assert result["status"] == "fail"
    failures = result["rule_audits"]["en"]["failures"]
    assert failures[0]["property_name"] == "Hotel Echo"
    assert failures[0]["expected_value"] == 187
    assert failures[0]["observed_value"] == 1_000_000


def test_standard_report_output_audit_blocks_implausible_airport_metric(
    tmp_path: Path,
) -> None:
    excel_path = tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx"
    _write_excel(excel_path)
    workbook = load_workbook(excel_path)
    workbook["Recommendations"]["C2"] = "43,712,000,000"
    workbook.save(excel_path)
    _write_ppt(
        tmp_path / "isite_test_standard_ppt_en_20260707T000000.pptx",
        airport_metric="Annual passenger throughput: 43,712,000,000",
        hotel_metric="Rooms / keys: 187",
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
    )

    assert result["status"] == "fail"
    failures = result["rule_audits"]["en"]["failures"]
    sanity_failure = next(
        item for item in failures
        if item["issue"] == "excel_airport_metric_implausible_high"
    )
    assert sanity_failure["property_name"] == "Airport Alpha"
    assert sanity_failure["observed_value"] == 43_712_000_000


def test_excel_only_audit_does_not_require_ppt_or_images(tmp_path: Path) -> None:
    _write_standard_excel(
        tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx",
        country="Testland",
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
        artifact_mode="excel-only",
    )

    assert result["status"] == "pass"
    assert result["rule_audits"]["en"]["main_table_row_count"] == 1


def test_excel_only_audit_blocks_country_scope_mismatch(tmp_path: Path) -> None:
    _write_standard_excel(
        tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx",
        country="Elsewhere",
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
        artifact_mode="excel-only",
    )

    assert result["status"] == "fail"
    assert any(
        failure["issue"] == "excel_country_scope_mismatch"
        for failure in result["rule_audits"]["en"]["failures"]
    )


def test_standard_report_output_audit_blocks_implausible_hotel_keys(
    tmp_path: Path,
) -> None:
    excel_path = tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx"
    _write_excel(excel_path)
    workbook = load_workbook(excel_path)
    workbook["Recommendations"]["C3"] = "12,500"
    workbook.save(excel_path)
    _write_ppt(
        tmp_path / "isite_test_standard_ppt_en_20260707T000000.pptx",
        airport_metric="Terminal capacity: 4,200,000",
        hotel_metric="Rooms / keys: 12,500",
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
    )

    assert result["status"] == "fail"
    assert any(
        failure["issue"] == "excel_hotel_keys_implausible_high"
        for failure in result["rule_audits"]["en"]["failures"]
    )

def test_standard_report_output_audit_blocks_one_page_recommendation_total_mismatch(
    tmp_path: Path,
) -> None:
    _write_excel(tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx")
    _write_ppt(
        tmp_path / "isite_test_standard_ppt_en_20260707T000000.pptx",
        airport_metric="Terminal capacity: 4,200,000",
        hotel_metric="Rooms / keys: 187",
        recommendation_total=3,
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
    )

    assert result["status"] == "fail"
    failures = result["rule_audits"]["en"]["failures"]
    assert any(item["issue"] == "ppt_recommendation_total_mismatch" for item in failures)


def test_standard_report_output_audit_allows_excel_scenes_not_shown_on_one_page(
    tmp_path: Path,
) -> None:
    excel_path = tmp_path / "isite_test_standard_report_en_20260707T000000.xlsx"
    _write_excel(excel_path)
    workbook = load_workbook(excel_path)
    workbook["Recommendations"].append(["Port Foxtrot", "Cruise Port", "5"])
    workbook.save(excel_path)
    _write_ppt(
        tmp_path / "isite_test_standard_ppt_en_20260707T000000.pptx",
        airport_metric="Terminal capacity: 4,200,000",
        hotel_metric="Rooms / keys: 187",
        recommendation_total=3,
    )

    result = audit_report_dir(
        tmp_path,
        country="Testland",
        locales=("en",),
        gpt_provider="off",
    )

    assert result["status"] == "pass"
    assert result["rule_audits"]["en"]["ppt_scene_card_recommendation_total"] == 2


def test_batch_output_audit_calls_gpt_once_for_multiple_countries(
    tmp_path: Path,
    monkeypatch,
) -> None:
    report_dirs = {}
    for country in ("Alpha", "Beta"):
        report_dir = tmp_path / country.casefold()
        report_dir.mkdir()
        _write_excel(
            report_dir / f"isite_{country.casefold()}_standard_report_en_20260707.xlsx"
        )
        _write_ppt(
            report_dir / f"isite_{country.casefold()}_standard_ppt_en_20260707.pptx",
            airport_metric="Terminal capacity: 4,200,000",
            hotel_metric="Rooms / keys: 187",
        )
        report_dirs[country] = report_dir

    calls = []

    def fake_gpt_audit(request, **_kwargs):
        calls.append(request)
        return {
            "verdict": "pass",
            "blocking_issues": [],
            "warnings": [],
            "checked_items": ["batch"],
            "summary": "pass",
        }

    monkeypatch.setattr(report_audit, "_run_codex_gpt_audit", fake_gpt_audit)

    result = report_audit.audit_report_dirs(
        report_dirs,
        locales=("en",),
        gpt_provider="codex-oauth",
    )

    assert result["status"] == "pass"
    assert len(calls) == 1
    assert set(calls[0]["countries"]) == {"Alpha", "Beta"}


def test_report_content_hash_is_stable_and_changes_with_payload(
    tmp_path: Path,
    monkeypatch,
) -> None:
    contract = tmp_path / "contract.txt"
    contract.write_text("v1", encoding="utf-8")
    monkeypatch.setattr(
        report_generation,
        "_report_contract_paths",
        lambda _artifact_mode="all": [contract],
    )

    first = report_generation._report_content_hash({"country": "Alpha", "count": 1})
    second = report_generation._report_content_hash({"count": 1, "country": "Alpha"})
    changed = report_generation._report_content_hash({"country": "Alpha", "count": 2})

    assert first == second
    assert first != changed


def test_qa_lessons_path_falls_back_to_latest_archive(tmp_path: Path) -> None:
    archive = tmp_path / "docs" / "qa_archive"
    archive.mkdir(parents=True)
    older = archive / "2025_history.md"
    latest = archive / "2026_history.md"
    older.write_text("older", encoding="utf-8")
    latest.write_text("latest", encoding="utf-8")

    assert report_audit.resolve_qa_lessons_path(tmp_path) == latest

    canonical = tmp_path / "docs" / "15_qa_lessons_learned.md"
    canonical.write_text("current", encoding="utf-8")
    assert report_audit.resolve_qa_lessons_path(tmp_path) == canonical


def test_reusable_report_requires_matching_hash_and_complete_artifacts(
    tmp_path: Path,
) -> None:
    previous = tmp_path / "alpha_standard_report_20260701"
    previous.mkdir()
    timestamp = "20260701"
    for path in report_generation._artifact_paths_for(
        previous,
        "alpha",
        ("zh", "en"),
        timestamp,
    ):
        path.write_bytes(b"artifact")
    (previous / f"standard_report_generation_summary_{timestamp}.json").write_text(
        (
            '{"timestamp":"20260701","content_hash":"same-hash",'
            '"audit_status":"pass"}'
        ),
        encoding="utf-8",
    )
    job = SimpleNamespace(
        slug="alpha",
        report_dir=tmp_path / "alpha_standard_report_20260702",
        content_hash="same-hash",
    )

    reusable = report_generation._find_reusable_report(
        tmp_path,
        job,
        ("zh", "en"),
    )

    assert reusable == previous


def test_excel_only_artifact_paths_exclude_ppt(tmp_path: Path) -> None:
    paths = report_generation._artifact_paths_for(
        tmp_path,
        "alpha",
        ("zh", "en"),
        "20260701",
        artifact_mode="excel-only",
    )

    assert len(paths) == 2
    assert all(path.suffix == ".xlsx" for path in paths)


def test_production_reuse_requires_passing_gpt_audit(tmp_path: Path) -> None:
    previous = tmp_path / "alpha_standard_report_20260701"
    previous.mkdir()
    timestamp = "20260701"
    for path in report_generation._artifact_paths_for(
        previous,
        "alpha",
        ("zh", "en"),
        timestamp,
    ):
        path.write_bytes(b"artifact")
    (previous / f"standard_report_generation_summary_{timestamp}.json").write_text(
        (
            '{"timestamp":"20260701","content_hash":"same-hash",'
            '"audit_status":"pass"}'
        ),
        encoding="utf-8",
    )
    job = SimpleNamespace(
        slug="alpha",
        report_dir=tmp_path / "alpha_standard_report_20260702",
        content_hash="same-hash",
    )

    assert (
        report_generation._find_reusable_report(
            tmp_path,
            job,
            ("zh", "en"),
            require_gpt_audit=True,
        )
        is None
    )
    (previous / f"standard_report_output_audit_{timestamp}.json").write_text(
        '{"status":"pass","gpt_audit":{"verdict":"pass"}}',
        encoding="utf-8",
    )

    assert (
        report_generation._find_reusable_report(
            tmp_path,
            job,
            ("zh", "en"),
            require_gpt_audit=True,
        )
        == previous
    )


def test_image_fallback_only_expands_scenes_still_missing_assets(tmp_path: Path) -> None:
    airport = SimpleNamespace(
        entity=SimpleNamespace(
            scene_type="airport_terminal",
            hero_image=SimpleNamespace(url="https://example.com/airport.jpg"),
        )
    )
    hotel_primary = SimpleNamespace(
        entity=SimpleNamespace(
            scene_type="luxury_hotel_mice",
            hero_image=SimpleNamespace(url="https://example.com/hotel-primary.jpg"),
        )
    )
    hotel_fallback = SimpleNamespace(
        entity=SimpleNamespace(
            scene_type="luxury_hotel_mice",
            hero_image=SimpleNamespace(url="https://example.com/hotel-fallback.jpg"),
        )
    )
    job = SimpleNamespace(
        export_packets=[airport, hotel_primary, hotel_fallback],
    )

    candidates = report_generation._missing_scene_image_candidates(
        job,
        {"https://example.com/airport.jpg": tmp_path / "airport.jpg"},
        {
            "https://example.com/airport.jpg",
            "https://example.com/hotel-primary.jpg",
        },
    )

    assert candidates == [hotel_fallback]


def _write_excel(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recommendations"
    sheet.append(
        [
            "Property Name",
            "Scene Type",
            "Threshold Metric Value",
            "Threshold Metric",
        ]
    )
    sheet.append(
        ["Airport Alpha", "Airport", "4,200,000", "annual_passenger_throughput"]
    )
    sheet.append(["Hotel Echo", "Luxury Hotel / MICE", "187", "keys"])
    workbook.save(path)


def _write_standard_excel(path: Path, *, country: str) -> None:
    labels = excel_labels("en")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for key in labels["required_sheet_keys"]:
        worksheet = workbook.create_sheet(labels["sheets"][key])
        if key == "main":
            worksheet.append(labels["columns"]["main"])
            row = [None] * len(labels["columns"]["main"])
            row[0] = country
            row[1] = "Test City"
            row[2] = "Airport Alpha"
            worksheet.append(row)
        elif key == "recommendation":
            worksheet.append(labels["columns"]["recommendation"])
            headers = labels["columns"]["recommendation"]
            row = [None] * len(headers)
            row[headers.index("Country")] = country
            row[headers.index("Property Name")] = "Airport Alpha"
            row[headers.index("Scene Type")] = "Airport"
            row[headers.index("Threshold Metric Value")] = "43,712,000"
            worksheet.append(row)
        else:
            worksheet.append(["Item"])
    workbook.save(path)


def _write_ppt(
    path: Path,
    *,
    airport_metric: str,
    hotel_metric: str,
    recommendation_total: int = 2,
) -> None:
    deck = Presentation()
    blank = deck.slide_layouts[6]
    slide = deck.slides.add_slide(blank)
    values = {
        "subtitle": f"2 qualified candidates · {recommendation_total} threshold recommendations · 2 scenes",
        "scene-01": "Airport",
        "counts-01": "Cand. 1 · Rec. 1",
        "name-01": "Airport Alpha",
        "metric-01": airport_metric,
        "status-01": "QUALIFIED",
        "scene-02": "Luxury Hotel / MICE",
        "counts-02": "Cand. 1 · Rec. 1",
        "name-02": "Hotel Echo",
        "metric-02": hotel_metric,
        "status-02": "QUALIFIED",
    }
    for index, (name, value) in enumerate(values.items()):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2 + index * 0.3), Inches(8), Inches(0.25))
        box.text = value
        box.name = name
    deck.save(path)
