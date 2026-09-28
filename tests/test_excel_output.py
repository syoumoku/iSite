from uuid import uuid4

import pytest
from openpyxl import load_workbook

from isite2.domain.enums import (
    ActionClass,
    BuildEvidenceStatus,
    CrossCheckStatus,
    EvidenceStatus,
    EvidenceType,
    IndoorRAT,
    IndoorSystemPresence,
    IndoorSystemType,
    ProxyLevel,
    RecommendedSolution,
    SceneForm,
    SourceTier,
    ValueClass,
)
from isite2.domain.models import (
    BuildStatus,
    Conclusion,
    EvidenceItem,
    PropertyEntity,
    SceneModelResult,
    SitePacket,
)
from isite2.localization import excel_labels, metric_label
from isite2.orchestrator.pipeline import run_scan_pipeline
from isite2.output.excel import (
    _metric_sort_value,
    create_excel_workbook,
    write_excel_skeleton,
)
from isite2.repositories.memory import InMemoryScanRunRepository
from isite2.rules.config_loader import scene_definitions


def test_excel_skeleton_matches_config_headers(tmp_path) -> None:
    output = tmp_path / "isite2_excel_skeleton.xlsx"
    write_excel_skeleton(output)

    workbook = load_workbook(output)
    excel = excel_labels("en")

    assert workbook.sheetnames == [excel["sheets"][key] for key in excel["required_sheet_keys"]]
    assert [cell.value for cell in workbook["Main"][1]] == excel["columns"]["main"]
    assert [cell.value for cell in workbook["Recommendations"][1]] == excel["columns"][
        "recommendation"
    ]
    assert [cell.value for cell in workbook["Airports"][1]] == excel["columns"]["main"]
    main_headers = [cell.value for cell in workbook["Main"][1]]
    airport_headers = [cell.value for cell in workbook["Airports"][1]]
    assert "Google Maps Link" in main_headers
    assert main_headers.index("Primary Metric Value") == (
        main_headers.index("Core Objective Evidence") + 1
    )
    assert airport_headers.index("Primary Metric Value") == (
        airport_headers.index("Core Objective Evidence") + 1
    )
    assert workbook["Proxy Model"].max_row == len(scene_definitions()) + 1
    assert workbook["Method"].max_row > 1
    assert any(
        row[0].value == "Recommendation Threshold Rule" for row in workbook["Method"].iter_rows()
    )
    assert "Recommended" in [cell.value for cell in workbook["Main"][1]]


def test_excel_skeleton_supports_chinese_locale(tmp_path) -> None:
    output = tmp_path / "isite2_excel_skeleton_zh.xlsx"
    write_excel_skeleton(output, locale="zh")

    workbook = load_workbook(output)
    excel = excel_labels("zh")

    assert workbook.sheetnames == [excel["sheets"][key] for key in excel["required_sheet_keys"]]
    assert [cell.value for cell in workbook["主表"][1]] == excel["columns"]["main"]
    main_headers = [cell.value for cell in workbook["主表"][1]]
    assert main_headers.index("主指标量化值") == main_headers.index("物业点重要证据") + 1
    assert "原始证据" in [cell.value for cell in workbook["证据表"][1]]
    assert "证据翻译" in [cell.value for cell in workbook["证据表"][1]]


def test_recommendation_metric_parser_does_not_double_scale_decimal_millions() -> None:
    assert _metric_sort_value(
        "annual_passenger_throughput: 43.712 million passengers in 2024",
        "annual_passenger_throughput",
    ) == 43_712_000
    assert _metric_sort_value(
        "annual_passenger_throughput: 16.095 million passengers",
        "annual_passenger_throughput",
    ) == pytest.approx(16_095_000)
    assert _metric_sort_value(
        "seat capacity: 15.000 seats",
        "seat_count",
    ) == 15_000


@pytest.mark.parametrize("locale", ["en", "zh"])
def test_excel_layout_preserves_full_names_and_recommendation_fill(locale) -> None:
    name = "Ulaanbaatar International Exhibition and Convention Centre " * 2
    packet = _packet_with_ordered_evidence(
        scene_type="stadium", weak_metric="event_role", strong_metric="seat_count",
        strong_value="30,000 seats", property_name=name,
    )
    workbook = create_excel_workbook([packet], locale=locale)
    labels = excel_labels(locale)
    for key in ("main", "recommendation", "stadium"):
        sheet = workbook[labels["sheets"][key]]
        assert sheet["C2"].value == name
        assert 14 <= sheet.column_dimensions["C"].width <= 60
        assert sheet["C2"].alignment.wrap_text
        assert sheet.row_dimensions[2].height > 30
        assert _is_recommended_fill(sheet["C2"])
        assert sheet.freeze_panes == "A2"


def test_excel_can_include_pipeline_packets(tmp_path) -> None:
    repository = InMemoryScanRunRepository()
    result = run_scan_pipeline(
        {
            "level": "country",
            "countries": ["Algeria"],
            "full_scan": True,
            "scene_types": ["airport_terminal"],
            "output_formats": ["excel"],
        },
        repository,
    )
    output = tmp_path / "isite2_excel_with_packets.xlsx"
    write_excel_skeleton(output, packets=result.packets)

    workbook = load_workbook(output)
    main_row = [cell.value for cell in workbook["Main"][2]]
    review_row = [cell.value for cell in workbook["Review Queue"][2]]
    google_maps_col = _header_index(workbook["Main"], "Google Maps Link")
    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")

    assert main_row[google_maps_col].startswith("https://www.google.com/maps/search/")
    assert primary_metric_col == _header_index(workbook["Main"], "Core Objective Evidence") + 1
    assert "Localization pending" not in main_row[6]
    assert not any("\u4e00" <= character <= "\u9fff" for character in main_row[6])
    assert "Localization pending" not in review_row[-1]


def test_excel_excludes_blocked_quality_candidates_by_default(tmp_path) -> None:
    repository = InMemoryScanRunRepository()
    result = run_scan_pipeline(
        {
            "level": "country",
            "countries": ["Exampleland"],
            "full_scan": True,
            "scene_types": ["airport_terminal"],
            "output_formats": ["excel"],
        },
        repository,
    )
    output = tmp_path / "isite2_excel_blocked_filtered.xlsx"

    write_excel_skeleton(output, packets=result.packets)

    workbook = load_workbook(output)
    assert workbook["Main"].max_row == 1


@pytest.mark.parametrize(
    ("scene_type", "weak_metric", "strong_metric", "strong_value", "expected_recommended"),
    [
        (
            "airport_terminal",
            "gateway_role",
            "annual_passenger_throughput",
            "GCAA 2024: 2,805,347 passengers/year.",
            True,
        ),
        (
            "convention_center",
            "venue_role",
            "exhibition_area",
            "20,000 sqm exhibition area.",
            False,
        ),
        ("stadium", "main_venue_role", "seat_count", "40,000 seats.", True),
        ("luxury_hotel_mice", "brand", "keys", "250 rooms.", True),
        ("mall_mixed_use", "flagship_position", "gla", "60,000 sqm GLA.", False),
        ("office_government", "cbd_role", "office_gfa", "80,000 sqm office GFA.", False),
        ("hospital", "referral_role", "beds", "600 beds.", True),
        ("university", "campus_center_role", "enrollment", "25,000 students.", True),
        (
            "transport_hub",
            "hub_role",
            "daily_ridership",
            "120,000 riders/day.",
            False,
        ),
    ],
)
def test_excel_main_row_prefers_scene_objective_metric(
    scene_type: str,
    weak_metric: str,
    strong_metric: str,
    strong_value: str,
    expected_recommended: bool,
) -> None:
    packet = _packet_with_ordered_evidence(
        scene_type=scene_type,
        weak_metric=weak_metric,
        strong_metric=strong_metric,
        strong_value=strong_value,
    )

    workbook = create_excel_workbook([packet])
    main_row = [cell.value for cell in workbook["Main"][2]]

    assert main_row[6].startswith(f"{metric_label(strong_metric, 'en')}:")
    assert strong_value in main_row[6]
    assert weak_metric not in main_row[6]
    assert main_row[_header_index(workbook["Main"], "Primary Metric Value")] is not None
    recommended_col = _header_index(workbook["Main"], "Recommended")
    reason_col = _header_index(workbook["Main"], "Recommendation Reason")
    assert (main_row[recommended_col] == "yes") is expected_recommended
    assert main_row[reason_col] == "Generic stale role-based reason."


def test_excel_main_table_sorts_by_annual_visits_desc() -> None:
    low = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Low Traffic Airport",
        annual_visits=1_200_000,
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value="1,200,000 passengers/year.",
    )
    missing = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Missing Traffic Airport",
        annual_visits=None,
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value="800,000 passengers/year.",
    )
    high = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="High Traffic Airport",
        annual_visits=9_800_000,
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value="9,800,000 passengers/year.",
    )

    workbook = create_excel_workbook([low, missing, high])

    assert [workbook["Main"][row][2].value for row in range(2, 5)] == [
        "High Traffic Airport",
        "Low Traffic Airport",
        "Missing Traffic Airport",
    ]


def test_excel_main_table_exposes_primary_metric_numeric_value() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Sortable Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value="annual_passenger_throughput: 92.3 million passengers in 2024",
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None)

    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][1][primary_metric_col].value == "Primary Metric Value"
    assert workbook["Main"][2][primary_metric_col].value == 92_300_000
    assert primary_metric_col == _header_index(workbook["Main"], "Core Objective Evidence") + 1

    airport_metric_col = _header_index(workbook["Airports"], "Primary Metric Value")
    assert airport_metric_col == _header_index(workbook["Airports"], "Core Objective Evidence") + 1
    assert workbook["Airports"][2][airport_metric_col].value == 92_300_000


def test_excel_recommendation_prefers_current_tier_one_metric_over_larger_weak_sources() -> None:
    hospital = _packet_with_ordered_evidence(
        scene_type="hospital",
        property_name="Current Evidence Hospital",
        weak_metric="referral_role",
        strong_metric="beds",
        strong_value="Current official profile: 500-bed facility.",
    )
    hospital.evidence[1].source_date = "2026-07-15"
    hospital.evidence.extend(
        [
            EvidenceItem(
                property_id=hospital.entity.property_id,
                field_group="beds",
                indicator_name="beds",
                field_value="Older secondary source: 340 beds.",
                source_name="Older secondary source",
                source_tier=SourceTier.TIER_2,
                source_url="https://example.com/older-secondary",
                source_date="2011",
                evidence_type=EvidenceType.DIRECT,
                cross_check_status=CrossCheckStatus.PARTIAL,
            ),
            EvidenceItem(
                property_id=hospital.entity.property_id,
                field_group="beds",
                indicator_name="beds",
                field_value="Undated directory: 700 beds.",
                source_name="Undated directory",
                source_tier=SourceTier.TIER_3,
                source_url="https://example.com/undated-directory",
                evidence_type=EvidenceType.DIRECT,
                cross_check_status=CrossCheckStatus.SINGLE_SOURCE,
            ),
        ]
    )

    workbook = create_excel_workbook([hospital], metric_identity_provider=None)

    metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][2][metric_col].value == 500
    assert workbook["Recommendations"][2][7].value == "500"


def test_excel_airport_primary_metric_value_skips_reporting_year() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Year Prefix Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value=(
            "annual_passenger_throughput: 2024 total passengers 12.9M; "
            "new terminal first-stage capacity 20 million passengers annually"
        ),
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None)

    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][2][primary_metric_col].value == 12_900_000
    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet[2][7].value == "12,900,000"


def test_excel_airport_primary_metric_accepts_served_people_in_airport_context() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Almaty Style Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value=(
            "annual_passenger_throughput: In 2023, the airport served "
            "9.5 million people."
        ),
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None)

    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][2][primary_metric_col].value == 9_500_000
    assert workbook["Recommendations"][2][7].value == "9,500,000"


def test_excel_airport_primary_metric_value_sums_domestic_and_international() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Split Traffic Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value=(
            "annual_passenger_throughput: 2025 international passenger numbers 678,591 "
            "plus domestic passenger numbers 1,023,529"
        ),
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None)

    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][2][primary_metric_col].value == 1_702_120
    assert workbook["Recommendations"].max_row == 1


def test_excel_university_primary_metric_value_skips_enrollment_year() -> None:
    university = _packet_with_ordered_evidence(
        scene_type="university",
        property_name="Enrollment Year University",
        weak_metric="campus_center_role",
        strong_metric="enrollment",
        strong_value="Enrollment: 870 students in 2023",
    )

    workbook = create_excel_workbook([university], metric_identity_provider=None)

    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][2][primary_metric_col].value == 870
    university_metric_col = _header_index(workbook["Universities"], "Primary Metric Value")
    assert workbook["Universities"][2][university_metric_col].value == 870


def test_excel_main_table_leaves_primary_metric_value_blank_for_descriptive_evidence() -> None:
    packet = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Role Only Airport",
        weak_metric="gateway_role",
        strong_metric="gateway_role",
        strong_value="Gateway airport opened in 2024 with no passenger throughput disclosed.",
    )

    workbook = create_excel_workbook([packet], metric_identity_provider=None)

    primary_metric_col = _header_index(workbook["Main"], "Primary Metric Value")
    assert workbook["Main"][2][primary_metric_col].value is None


def test_excel_scene_sheet_sorts_by_scene_primary_metric_desc() -> None:
    medium = _packet_with_ordered_evidence(
        scene_type="stadium",
        property_name="Medium Stadium",
        annual_visits=10_000_000,
        weak_metric="main_venue_role",
        strong_metric="seat_count",
        strong_value="45,000 seats.",
    )
    low = _packet_with_ordered_evidence(
        scene_type="stadium",
        property_name="Low Stadium",
        annual_visits=99_000_000,
        weak_metric="main_venue_role",
        strong_metric="seat_count",
        strong_value="20,000 seats.",
    )
    high = _packet_with_ordered_evidence(
        scene_type="stadium",
        property_name="High Stadium",
        annual_visits=1_000_000,
        weak_metric="main_venue_role",
        strong_metric="seat_count",
        strong_value="65,000 seats.",
    )

    workbook = create_excel_workbook([medium, low, high])

    assert [workbook["Stadiums"][row][2].value for row in range(2, 5)] == [
        "High Stadium",
        "Medium Stadium",
        "Low Stadium",
    ]


def test_excel_recommendation_sheet_and_highlights_threshold_rows() -> None:
    packets = [
        _packet_with_ordered_evidence(
            scene_type="stadium",
            property_name=f"Stadium {index:02d}",
            annual_visits=1_000_000,
            weak_metric="main_venue_role",
            strong_metric="seat_count",
            strong_value=f"{10_000 + index * 1_000:,} seats.",
        )
        for index in range(12)
    ]

    workbook = create_excel_workbook(packets)

    recommendation_sheet = workbook["Recommendations"]
    recommendation_names = [
        recommendation_sheet[row][2].value for row in range(2, recommendation_sheet.max_row + 1)
    ]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][5].value == "Seats > 20,000"
    assert recommendation_names == ["Stadium 11"]

    stadium_sheet = workbook["Stadiums"]
    recommended_col = _header_index(stadium_sheet, "Recommended")
    assert _is_recommended_fill(stadium_sheet[2][0])
    assert stadium_sheet[2][recommended_col].value == "yes"
    assert not _is_recommended_fill(stadium_sheet[3][0])
    assert stadium_sheet[3][recommended_col].value in {None, ""}

    main_rows = {row[2].value: row for row in workbook["Main"].iter_rows(min_row=2)}
    main_recommended_col = _header_index(workbook["Main"], "Recommended")
    assert _is_recommended_fill(main_rows["Stadium 11"][0])
    assert main_rows["Stadium 11"][main_recommended_col].value == "yes"
    assert not _is_recommended_fill(main_rows["Stadium 10"][0])
    assert main_rows["Stadium 10"][main_recommended_col].value in {None, ""}
    assert not _is_recommended_fill(main_rows["Stadium 00"][0])
    assert main_rows["Stadium 00"][main_recommended_col].value in {None, ""}


def test_excel_mall_recommendation_recognizes_retail_gfa_abl_context() -> None:
    mall = _packet_with_ordered_evidence(
        scene_type="mall_mixed_use",
        property_name="Large ABL Mall",
        weak_metric="flagship_position",
        strong_metric="retail_gfa",
        strong_value=(
            "retail gross floor area (ABL) in m2: "
            "retail gross floor area (ABL): 145000 m2"
        ),
    )

    workbook = create_excel_workbook([mall], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][2].value == "Large ABL Mall"
    assert recommendation_sheet[2][5].value == "GLA > 120,000 sqm"
    assert recommendation_sheet[2][6].value == "gla"
    assert recommendation_sheet[2][7].value == "145,000"
    mall_sheet = workbook["Mall Mixed-use"]
    recommended_col = _header_index(mall_sheet, "Recommended")
    assert mall_sheet[2][recommended_col].value == "yes"
    assert _is_recommended_fill(mall_sheet[2][0])


def test_excel_fixed_transport_gate_uses_line_count() -> None:
    single_line = _packet_with_ordered_evidence(
        scene_type="transport_hub",
        property_name="Single Line Station",
        weak_metric="hub_role",
        strong_metric="line_count",
        strong_value="1 metro line.",
    )
    interchange = _packet_with_ordered_evidence(
        scene_type="transport_hub",
        property_name="Interchange Station",
        weak_metric="hub_role",
        strong_metric="line_count",
        strong_value="2 metro lines.",
    )

    workbook = create_excel_workbook([single_line, interchange])

    recommendation_sheet = workbook["Recommendations"]
    recommendation_names = [
        recommendation_sheet[row][2].value for row in range(2, recommendation_sheet.max_row + 1)
    ]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][5].value == "Interchange Lines > 1"
    assert recommendation_names == ["Interchange Station"]


def test_excel_university_fixed_gate_includes_exactly_twenty_thousand() -> None:
    below_gate = _packet_with_ordered_evidence(
        scene_type="university",
        property_name="Below Gate University",
        weak_metric="campus_center_role",
        strong_metric="enrollment",
        strong_value="Enrollment: 19,999 students.",
    )
    at_gate = _packet_with_ordered_evidence(
        scene_type="university",
        property_name="At Gate University",
        weak_metric="campus_center_role",
        strong_metric="enrollment",
        strong_value="Enrollment: 20,000 students.",
    )

    workbook = create_excel_workbook([below_gate, at_gate], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][2].value == "At Gate University"
    assert recommendation_sheet[2][5].value == "Single-campus enrollment >= 20,000"
    assert recommendation_sheet[2][6].value == "enrollment"
    assert recommendation_sheet[2][7].value == "20,000"

    university_sheet = workbook["Universities"]
    rows = {row[2].value: row for row in university_sheet.iter_rows(min_row=2)}
    recommended_col = _header_index(university_sheet, "Recommended")
    assert rows["At Gate University"][recommended_col].value == "yes"
    assert rows["Below Gate University"][recommended_col].value in {None, ""}


def test_excel_hotel_room_gate_does_not_treat_conference_room_area_as_keys() -> None:
    hotel = _packet_with_ordered_evidence(
        scene_type="luxury_hotel_mice",
        property_name="Conference Area Hotel",
        weak_metric="brand",
        strong_metric="keys",
        strong_value=(
            "Northstar lists 539 guest rooms, 36 meeting rooms, "
            "largest conference room 10,368 sq. ft."
        ),
    )

    workbook = create_excel_workbook([hotel])

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][6].value == "keys"
    assert recommendation_sheet[2][7].value == "539"


def test_excel_fixed_convention_and_office_gates_use_area_and_height() -> None:
    convention_at_gate = _packet_with_ordered_evidence(
        scene_type="convention_center",
        property_name="At Gate Convention Center",
        weak_metric="venue_role",
        strong_metric="exhibition_area",
        strong_value="25,000 sqm exhibition area.",
    )
    convention_above_gate = _packet_with_ordered_evidence(
        scene_type="convention_center",
        property_name="Above Gate Convention Center",
        weak_metric="venue_role",
        strong_metric="exhibition_area",
        strong_value="26,000 sqm exhibition area.",
    )
    office_at_gate = _packet_with_ordered_evidence(
        scene_type="office_government",
        property_name="At Gate Office Tower",
        weak_metric="cbd_role",
        strong_metric="tower_height",
        strong_value="150 m tower height.",
    )
    office_above_gate = _packet_with_ordered_evidence(
        scene_type="office_government",
        property_name="Above Gate Office Tower",
        weak_metric="cbd_role",
        strong_metric="tower_height",
        strong_value="151 m tower height.",
    )

    workbook = create_excel_workbook(
        [
            convention_at_gate,
            convention_above_gate,
            office_at_gate,
            office_above_gate,
        ]
    )

    recommendation_sheet = workbook["Recommendations"]
    recommendation_rows = {
        recommendation_sheet[row][2].value: recommendation_sheet[row]
        for row in range(2, recommendation_sheet.max_row + 1)
    }
    assert set(recommendation_rows) == {
        "Above Gate Convention Center",
        "Above Gate Office Tower",
    }
    assert recommendation_rows["Above Gate Convention Center"][5].value == (
        "Exhibition Area > 25,000 sqm"
    )
    assert recommendation_rows["Above Gate Office Tower"][5].value == "Building Height > 150 m"

    main_rows = {row[2].value: row for row in workbook["Main"].iter_rows(min_row=2)}
    main_recommended_col = _header_index(workbook["Main"], "Recommended")
    assert main_rows["Above Gate Convention Center"][main_recommended_col].value == "yes"
    assert main_rows["At Gate Convention Center"][main_recommended_col].value in {None, ""}
    assert main_rows["Above Gate Office Tower"][main_recommended_col].value == "yes"
    assert main_rows["At Gate Office Tower"][main_recommended_col].value in {None, ""}


def test_excel_mall_fixed_gate_treats_retail_gfa_as_gla() -> None:
    below_gate = _packet_with_ordered_evidence(
        scene_type="mall_mixed_use",
        property_name="Below Gate Mall",
        weak_metric="flagship_position",
        strong_metric="retail_gfa",
        strong_value="Retail GFA: 90,000 sqm.",
    )
    above_gate = _packet_with_ordered_evidence(
        scene_type="mall_mixed_use",
        property_name="Above Gate Mall",
        weak_metric="flagship_position",
        strong_metric="retail_gfa",
        strong_value="Retail GFA: 170,000 sqm.",
    )

    workbook = create_excel_workbook([below_gate, above_gate])

    recommendation_sheet = workbook["Recommendations"]
    recommendation_rows = {
        recommendation_sheet[row][2].value: recommendation_sheet[row]
        for row in range(2, recommendation_sheet.max_row + 1)
    }
    assert set(recommendation_rows) == {"Above Gate Mall"}
    assert recommendation_rows["Above Gate Mall"][5].value == "GLA > 120,000 sqm"
    assert recommendation_rows["Above Gate Mall"][6].value == "gla"

    main_rows = {row[2].value: row for row in workbook["Main"].iter_rows(min_row=2)}
    recommended_col = _header_index(workbook["Main"], "Recommended")
    assert main_rows["Above Gate Mall"][recommended_col].value == "yes"
    assert main_rows["Below Gate Mall"][recommended_col].value in {None, ""}


def test_excel_mall_gla_gate_does_not_use_annual_footfall_as_gla() -> None:
    mall = _packet_with_ordered_evidence(
        scene_type="mall_mixed_use",
        property_name="Footfall Heavy Mall",
        weak_metric="flagship_position",
        strong_metric="retail_gfa",
        strong_value="GLA in SQM: 18,259; annual footfall: 7,741,200",
    )

    workbook = create_excel_workbook([mall], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 1

    main_rows = {row[2].value: row for row in workbook["Main"].iter_rows(min_row=2)}
    recommended_col = _header_index(workbook["Main"], "Recommended")
    assert main_rows["Footfall Heavy Mall"][recommended_col].value in {None, ""}


def test_excel_office_fixed_gate_accepts_height_unit_in_label() -> None:
    office = _packet_with_ordered_evidence(
        scene_type="office_government",
        property_name="Height Label Office Tower",
        weak_metric="cbd_role",
        strong_metric="tower_height",
        strong_value="Height (m): 192 meters",
    )

    workbook = create_excel_workbook([office], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][2].value == "Height Label Office Tower"
    assert recommendation_sheet[2][6].value == "tower_height"
    assert recommendation_sheet[2][7].value == "192"


def test_excel_stadium_fixed_gate_accepts_number_before_capacity_label() -> None:
    stadium = _packet_with_ordered_evidence(
        scene_type="stadium",
        property_name="Number Before Capacity Stadium",
        weak_metric="main_venue_role",
        strong_metric="seat_count",
        strong_value="seat_count: 40,000 capacity",
    )

    workbook = create_excel_workbook([stadium], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][2].value == "Number Before Capacity Stadium"
    assert recommendation_sheet[2][6].value == "seat_count"
    assert recommendation_sheet[2][7].value == "40,000"


def test_excel_fixed_gate_uses_canonical_metric_prefix_in_field_value() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Canonical Prefix Airport",
        weak_metric="terminal_capacity",
        strong_metric="terminal_capacity",
        strong_value=(
            "annual_passenger_throughput: airport handled 2,993,453 passengers "
            "and has annual passenger capacity over 5 million."
        ),
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][2].value == "Canonical Prefix Airport"
    assert recommendation_sheet[2][6].value == "annual_passenger_throughput"
    assert recommendation_sheet[2][7].value == "2,993,453"


def test_excel_airport_gate_accepts_million_guests_passengers_label() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Guests Slash Passengers Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value="annual_passenger_throughput: 92.3 million guests/passengers in 2024",
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None)

    recommendation_sheet = workbook["Recommendations"]
    assert recommendation_sheet.max_row == 2
    assert recommendation_sheet[2][2].value == "Guests Slash Passengers Airport"
    assert recommendation_sheet[2][6].value == "annual_passenger_throughput"
    assert recommendation_sheet[2][7].value == "92,300,000"


def test_excel_airport_gate_accepts_terminal_capacity_passengers_per_year() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Annual Terminal Capacity Airport",
        weak_metric="gateway_role",
        strong_metric="terminal_capacity",
        strong_value=(
            "terminal_capacity: 3,000,000 passengers per year planned handling capacity; "
            "new passenger terminal of 32,000 m2"
        ),
    )
    small_airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Hourly Capacity Airport",
        weak_metric="gateway_role",
        strong_metric="terminal_capacity",
        strong_value="terminal capacity: 500 passengers per hour",
    )

    workbook = create_excel_workbook(
        [airport, small_airport],
        metric_identity_provider=None,
    )

    recommendation_sheet = workbook["Recommendations"]
    recommendation_names = [
        recommendation_sheet[row][2].value for row in range(2, recommendation_sheet.max_row + 1)
    ]
    assert recommendation_names == ["Annual Terminal Capacity Airport"]
    assert recommendation_sheet[2][6].value == "terminal_capacity"
    assert recommendation_sheet[2][7].value == "3,000,000"


def test_excel_hotel_gate_ignores_implausible_room_count_evidence() -> None:
    hotel = _packet_with_ordered_evidence(
        scene_type="luxury_hotel_mice",
        property_name="President Hotel",
        weak_metric="brand",
        strong_metric="keys",
        strong_value="keys: 204 rooms from OSM rooms tag",
    )
    hotel.evidence[-1].source_tier = SourceTier.TIER_3
    hotel.evidence.append(
        EvidenceItem(
            property_id=hotel.entity.property_id,
            field_group="keys",
            indicator_name="keys",
            field_value="12,500 rooms",
            source_name="Misclassified official-page number",
            source_tier=SourceTier.TIER_1,
            source_url="https://example.com/unrelated-number",
            evidence_type=EvidenceType.DIRECT,
            cross_check_status=CrossCheckStatus.SINGLE_SOURCE,
        )
    )

    workbook = create_excel_workbook([hotel], metric_identity_provider=None)

    recommendation = workbook["Recommendations"][2]
    assert recommendation[2].value == "President Hotel"
    assert recommendation[6].value == "keys"
    assert recommendation[7].value == "204"


def test_excel_english_output_removes_chinese_template_residue() -> None:
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="Localized Fallback Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_throughput",
        strong_value="主指标：annual_passenger_throughput: 4.9 million passengers in 2025",
    )
    airport.conclusion.reason_to_recommend = (
        "门户机场, January-May 2025 passenger movement: 4.9 million passengers. "
        "Terminal indoor continuous coverage demand is clear, recommend pRRU."
    )

    workbook = create_excel_workbook([airport], metric_identity_provider=None, locale="en")
    values = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    values.append(cell.value)

    assert not any(
        any("\u4e00" <= character <= "\u9fff" for character in value)
        for value in values
    )
    assert workbook["Main"][2][6].value.startswith(
        "Annual passenger throughput: Primary metric:"
    )
    evidence_values = [
        workbook["Evidence"][row][6].value
        for row in range(2, workbook["Evidence"].max_row + 1)
    ]
    assert any(
        str(value).startswith("Primary metric: annual_passenger_throughput:")
        for value in evidence_values
    )
    reason_col = _header_index(workbook["Main"], "Recommendation Reason")
    assert workbook["Main"][2][reason_col].value.startswith("Gateway airport,")
    assert "Localization pending" not in workbook["Main"][2][reason_col].value


def test_excel_recommendation_metric_gpt_normalizes_all_scene_gate_metrics() -> None:
    provider = _RecordingMetricIdentityProvider(
        {
            "annual_passenger_count": "annual_passenger_throughput",
            "venue_capacity": "seat_count",
        }
    )
    airport = _packet_with_ordered_evidence(
        scene_type="airport_terminal",
        property_name="GPT Airport",
        weak_metric="gateway_role",
        strong_metric="annual_passenger_count",
        strong_value="2024 annual passenger count: 3,100,000 passengers.",
    )
    stadium = _packet_with_ordered_evidence(
        scene_type="stadium",
        property_name="GPT Stadium",
        weak_metric="main_venue_role",
        strong_metric="venue_capacity",
        strong_value="Venue capacity is 35,000 seats.",
    )

    workbook = create_excel_workbook(
        [airport, stadium],
        metric_identity_provider=provider,
    )

    recommendation_sheet = workbook["Recommendations"]
    recommendation_rows = {
        recommendation_sheet[row][2].value: recommendation_sheet[row]
        for row in range(2, recommendation_sheet.max_row + 1)
    }
    assert set(recommendation_rows) == {"GPT Airport", "GPT Stadium"}
    assert recommendation_rows["GPT Airport"][6].value == "annual_passenger_throughput"
    assert recommendation_rows["GPT Stadium"][6].value == "seat_count"
    assert {
        request["scene_type"]
        for request in provider.requests
        if request["raw_metric"]["field_group"] in {"annual_passenger_count", "venue_capacity"}
    } == {"airport_terminal", "stadium"}


class _RecordingMetricIdentityProvider:
    provider_name = "fake_gpt_metric_identity"

    def __init__(self, mapping: dict[str, str]) -> None:
        self.mapping = mapping
        self.requests: list[dict] = []

    def identify(self, request: dict) -> dict:
        self.requests.append(request)
        raw_key = request["raw_metric"]["field_group"]
        canonical = self.mapping.get(raw_key)
        return {
            "recognized": canonical is not None,
            "canonical_metric_key": canonical,
            "is_scene_primary_metric": canonical is not None,
            "confidence": 0.94 if canonical else 0.2,
            "reason": "fake test decision",
        }


def _packet_with_ordered_evidence(
    *,
    scene_type: str,
    weak_metric: str,
    strong_metric: str,
    strong_value: str,
    property_name: str = "Objective Metric Test Property",
    annual_visits: int | None = 1_000_000,
) -> SitePacket:
    property_id = uuid4()
    scene_rule = scene_definitions()[scene_type]
    scene_form = (
        SceneForm.SEMI_OPEN
        if scene_rule["scene_form"] == SceneForm.SEMI_OPEN.value
        else SceneForm.INDOOR
    )
    return SitePacket(
        entity=PropertyEntity(
            property_id=property_id,
            country="Ghana",
            city="Accra",
            property_name=property_name,
            scene_type=scene_type,
            scene_form=scene_form,
            latitude=5.6,
            longitude=-0.1,
            geocode_precision="Property",
            google_maps_link="https://www.google.com/maps/search/?api=1&query=5.6,-0.1",
        ),
        scene=SceneModelResult(
            area_metric_name=scene_rule["area_metric"],
            primary_value_indicators=list(scene_rule["primary_indicators"]),
            proxy_basis=scene_rule["proxy_basis"][0],
            proxy_level=ProxyLevel.P1_STRONG,
            annual_visits_est=annual_visits,
        ),
        evidence=[
            EvidenceItem(
                property_id=property_id,
                field_group=weak_metric,
                indicator_name=weak_metric,
                field_value="National role / gateway description only.",
                source_name="Directory seed",
                source_tier=SourceTier.TIER_3,
                source_url="https://example.com/role",
                evidence_type=EvidenceType.DIRECT,
                cross_check_status=CrossCheckStatus.SINGLE_SOURCE,
            ),
            EvidenceItem(
                property_id=property_id,
                field_group=strong_metric,
                indicator_name=strong_metric,
                field_value=strong_value,
                source_name="Official statistics",
                source_tier=SourceTier.TIER_1,
                source_url="https://example.com/statistics",
                evidence_type=EvidenceType.DIRECT,
                cross_check_status=CrossCheckStatus.PARTIAL,
            ),
        ],
        build_status=BuildStatus(
            indoor_system_presence=IndoorSystemPresence.NO_PUBLIC_EVIDENCE,
            indoor_system_type=IndoorSystemType.UNKNOWN,
            indoor_rat=IndoorRAT.UNKNOWN,
            build_evidence_status=BuildEvidenceStatus.UNKNOWN,
        ),
        conclusion=Conclusion(
            evidence_status=EvidenceStatus.SUPPORTED,
            value_class=ValueClass.CITY_CORE,
            action_class=ActionClass.SURVEY_FIRST,
            recommended_solution=RecommendedSolution.PRRU,
            reason_to_recommend="Generic stale role-based reason.",
            next_action="补查现网室分状态和运营商证据。",
        ),
    )


def _is_recommended_fill(cell) -> bool:
    return cell.fill.fgColor.rgb in {"00FFF2CC", "FFFFF2CC"}


def _header_index(worksheet, header: str) -> int:
    headers = [cell.value for cell in worksheet[1]]
    return headers.index(header)
