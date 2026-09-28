from __future__ import annotations

import concurrent.futures
import hashlib
import json
import os
import re
import shutil
from datetime import date
from pathlib import Path

import httpx
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from isite2.domain.models import SitePacket
from isite2.localization import (
    DatabaseLocalizationCache,
    enum_label,
    generic_free_text_fallback,
    localize_text,
    metric_label,
    ppt_text,
    resolve_locale,
    scene_label,
    t,
)
from isite2.output.excel import (
    RecommendationDecision,
    _build_recommendation_index,
    _evidence_metric_key,
    _format_metric_value,
    _main_metric_text,
    _metric_sort_value,
    _packet_key,
    _recommended_packets,
    _select_main_evidence,
)
from isite2.media.hero_images import hero_image_url_variants, stable_hero_image_url
from isite2.rules.candidate_quality import filter_packets_for_surface
from isite2.rules.config_loader import load_output_template, scene_definitions


_SCENE_ORDER = [
    "airport_terminal",
    "transport_hub",
    "mall_mixed_use",
    "stadium",
    "luxury_hotel_mice",
    "convention_center",
    "office_government",
    "hospital",
    "university",
    "mosque",
]
_METHODOLOGY_SCENES = [
    "airport_terminal",
    "transport_hub",
    "mall_mixed_use",
    "stadium",
    "luxury_hotel_mice",
    "convention_center",
    "office_government",
    "hospital",
]
_BLUE = RGBColor(31, 78, 121)
_LIGHT_BLUE = RGBColor(221, 235, 247)
_PALE_YELLOW = RGBColor(255, 242, 204)
_PALE_GREEN = RGBColor(226, 239, 218)
_GRAY = RGBColor(89, 89, 89)
_BORDER = RGBColor(190, 190, 190)
_ROOT = Path(__file__).resolve().parents[3]
_ONE_PAGE_TEMPLATES = {
    "zh": _ROOT / "templates" / "ppt" / "isite_standard_one_page_zh.pptx",
    "en": _ROOT / "templates" / "ppt" / "isite_standard_one_page_en.pptx",
}
_CARD_SCENE_LABELS = {
    "airport_terminal": {"zh": "机场", "en": "Airport"},
    "transport_hub": {"zh": "交通枢纽", "en": "Transport Hub"},
    "convention_center": {"zh": "会展中心", "en": "Convention"},
    "mall_mixed_use": {"zh": "大型商超", "en": "Large Retail"},
    "stadium": {"zh": "体育场", "en": "Stadium"},
    "luxury_hotel_mice": {"zh": "奢华酒店", "en": "Luxury Hotel"},
    "office_government": {"zh": "写字楼", "en": "Office Tower"},
    "hospital": {"zh": "医院", "en": "Hospital"},
    "university": {"zh": "大学", "en": "University"},
    "mosque": {"zh": "清真寺", "en": "Mosque"},
}


def scene_card_label(scene_type: str, locale: str) -> str:
    """Return the compact scene label used by the one-page PPT cards."""
    return (
        _CARD_SCENE_LABELS.get(scene_type, {}).get(locale)
        or scene_label(scene_type, locale)
    )


def create_ppt_deck(
    packets: list[SitePacket],
    title: str = "iSite2 Insight Cards",
    include_blocked_quality: bool = False,
    locale: str | None = None,
    localization_cache: DatabaseLocalizationCache | None = None,
    recommendations: dict[str, RecommendationDecision] | None = None,
    image_assets: dict[str, Path] | None = None,
    require_images: bool = False,
) -> Presentation:
    resolved_locale = resolve_locale(locale)
    packets = filter_packets_for_surface(
        packets,
        "export",
        include_blocked_quality=include_blocked_quality,
    )
    recommendations = recommendations or _build_recommendation_index(
        packets,
        locale=resolved_locale,
    )
    deck = Presentation(_ONE_PAGE_TEMPLATES[resolved_locale])
    _populate_one_page_slide(
        deck,
        packets,
        recommendations,
        resolved_locale,
        localization_cache=localization_cache,
        image_assets=image_assets or {},
        require_images=require_images,
    )
    return deck


def write_ppt_deck(
    path: Path,
    packets: list[SitePacket],
    title: str = "iSite2 Insight Cards",
    include_blocked_quality: bool = False,
    locale: str | None = None,
    localization_cache: DatabaseLocalizationCache | None = None,
    recommendations: dict[str, RecommendationDecision] | None = None,
    image_assets: dict[str, Path] | None = None,
    require_images: bool = False,
) -> Path:
    deck = create_ppt_deck(
        packets,
        title,
        include_blocked_quality,
        locale=locale,
        localization_cache=localization_cache,
        recommendations=recommendations,
        image_assets=image_assets,
        require_images=require_images,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)
    return path


def prepare_ppt_assets(
    packets: list[SitePacket],
    *,
    max_workers: int = 8,
) -> dict[str, Path]:
    """Resolve real-property hero images once for reuse by both locales."""
    tasks = {
        str(packet.entity.hero_image.url): packet.entity.property_name
        for packet in packets
        if packet.entity.hero_image and packet.entity.hero_image.url
    }
    if not tasks:
        return {}
    evidence_images = _local_evidence_image_index()
    assets: dict[str, Path] = {}
    with concurrent.futures.ThreadPoolExecutor(
        max_workers=max(1, min(max_workers, len(tasks)))
    ) as executor:
        future_by_url = {
            executor.submit(
                _resolve_ppt_image,
                url,
                _matching_local_evidence_image(
                    property_name,
                    evidence_images,
                ),
            ): url
            for url, property_name in tasks.items()
        }
        for future in concurrent.futures.as_completed(future_by_url):
            url = future_by_url[future]
            path = future.result()
            if path is not None:
                assets[url] = path
    return assets


def scene_card_packets(
    packets: list[SitePacket],
    recommendations: dict[str, RecommendationDecision],
    *,
    image_assets: dict[str, Path] | None = None,
    require_images: bool = False,
) -> list[SitePacket]:
    grouped: dict[str, list[SitePacket]] = {}
    for packet in packets:
        grouped.setdefault(packet.entity.scene_type, []).append(packet)
    scenes = scene_card_scene_types(packets)
    selected: list[SitePacket] = []
    for scene_type in scenes:
        scene_packets = grouped[scene_type]
        scene_packets.sort(
            key=lambda packet: _scene_card_sort_key(
                packet,
                recommendations.get(_packet_key(packet)),
                image_assets=image_assets,
            )
        )
        if require_images and not _packet_has_image_asset(
            scene_packets[0],
            image_assets or {},
        ):
            raise ValueError(
                f"missing cached real property image for scene {scene_type}"
            )
        selected.append(scene_packets[0])
    return selected


def scene_card_scene_types(
    packets: list[SitePacket],
    *,
    limit: int = 9,
) -> list[str]:
    """Choose display scenes without truncating the underlying export pool."""
    available = {packet.entity.scene_type for packet in packets}
    scenes = [scene for scene in _SCENE_ORDER if scene in available]
    scenes.extend(sorted(available - set(scenes)))
    return scenes[: max(0, limit)]


def scene_card_image_candidates(
    packets: list[SitePacket],
    recommendations: dict[str, RecommendationDecision],
    *,
    per_scene: int = 3,
) -> list[SitePacket]:
    """Return a small ranked image fallback pool for each one-page scene card."""
    grouped: dict[str, list[SitePacket]] = {}
    for packet in packets:
        grouped.setdefault(packet.entity.scene_type, []).append(packet)
    candidates: list[SitePacket] = []
    scenes = scene_card_scene_types(packets)
    for scene_type in scenes:
        scene_packets = grouped[scene_type]
        scene_packets.sort(
            key=lambda packet: _scene_card_sort_key(
                packet,
                recommendations.get(_packet_key(packet)),
            )
        )
        candidates.extend(scene_packets[: max(1, per_scene)])
    return candidates


def _populate_one_page_slide(
    deck: Presentation,
    packets: list[SitePacket],
    recommendations: dict[str, RecommendationDecision],
    locale: str,
    *,
    localization_cache: DatabaseLocalizationCache | None,
    image_assets: dict[str, Path],
    require_images: bool,
) -> None:
    if len(deck.slides) != 1:
        raise ValueError("one-page standard template must contain exactly one slide")
    slide = deck.slides[0]
    shapes = {shape.name: shape for shape in slide.shapes if shape.name}
    pictures = sorted(
        [shape for shape in slide.shapes if shape.shape_type == 13],
        key=lambda shape: (shape.top, shape.left),
    )
    representatives = scene_card_packets(
        packets,
        recommendations,
        image_assets=image_assets,
        require_images=require_images,
    )
    _reflow_scene_cards(slide, representatives, pictures)

    country = _country_names(packets)[0] if packets else ""
    country_title = _localized_country_name(country, locale, localization_cache)
    recommended_packets = _recommended_packets(packets, recommendations)
    scene_counts = _scene_counts(packets, recommendations)
    title = (
        f"{country_title}高价值楼宇洞察"
        if locale == "zh"
        else f"{country_title} High-Value Building Insights"
    )
    subtitle = (
        f"{len(packets)} 个合格候选  ·  {len(recommended_packets)} 个门槛推荐  ·  {len(scene_counts)} 类场景"
        if locale == "zh"
        else f"{len(packets)} qualified candidates  ·  {len(recommended_packets)} threshold recommendations  ·  {len(scene_counts)} scenes"
    )
    _set_shape_text(shapes["title"], title, size=_title_font_size(title, locale))
    _set_shape_text(shapes["subtitle"], subtitle)
    _set_shape_text(shapes["insight-distribution"], _recommendation_mix(scene_counts, locale))
    _set_shape_text(
        shapes["footer-note"],
        (
            f"场景卡融合候选规模、推荐数量、门槛、代表物业和同口径主指标；数值复核截至 {date.today().isoformat()}。"
            if locale == "zh"
            else f"Each scene card combines pool size, recommendations, threshold, representative site and same-basis hard metric. Audited {date.today().isoformat()}."
        ),
    )
    _apply_brand_flag(shapes)

    for index, packet in enumerate(representatives, start=1):
        suffix = f"{index:02d}"
        decision = recommendations[_packet_key(packet)]
        scene_type = packet.entity.scene_type
        candidate_count, recommended_count = scene_counts[scene_type]
        _set_shape_text(
            shapes[f"scene-{suffix}"],
            scene_card_label(scene_type, locale),
        )
        _set_shape_text(
            shapes[f"counts-{suffix}"],
            (
                f"候选 {candidate_count}  ·  推荐 {recommended_count}"
                if locale == "zh"
                else f"Cand. {candidate_count}  ·  Rec. {recommended_count}"
            ),
        )
        property_name = packet.entity.property_name
        _set_shape_text(
            shapes[f"name-{suffix}"],
            property_name,
            size=_property_font_size(property_name, locale),
        )
        metric_text = _card_metric_text(packet, decision, locale)
        _set_shape_text(
            shapes[f"metric-{suffix}"],
            metric_text,
            size=_card_metric_font_size(
                metric_text,
                locale=locale,
                dense=len(representatives) >= 9,
            ),
        )
        _set_shape_text(
            shapes[f"threshold-{suffix}"],
            _card_threshold_text(decision, locale),
            size=8.5 if len(representatives) >= 9 else 9.5,
        )
        status_text, status_fill, status_color = _card_status(decision, locale)
        _set_shape_text(
            shapes[f"status-{suffix}"],
            status_text,
            size=7.5 if len(representatives) >= 9 else 8.5,
        )
        _set_shape_fill(shapes[f"status-bg-{suffix}"], status_fill)
        _set_shape_font_color(shapes[f"status-{suffix}"], status_color)
        hero = packet.entity.hero_image
        image_path = image_assets.get(str(hero.url)) if hero else None
        if image_path is None and require_images:
            raise ValueError(
                f"missing cached real property image for {packet.entity.property_name}"
            )
        picture = pictures[index - 1]
        picture.name = f"image-{suffix}"
        if image_path is not None:
            _replace_picture(slide, picture, image_path)
        else:
            _remove_shape(picture)


def _scene_counts(
    packets: list[SitePacket],
    recommendations: dict[str, RecommendationDecision],
) -> dict[str, tuple[int, int]]:
    counts: dict[str, list[int]] = {}
    for packet in packets:
        item = counts.setdefault(packet.entity.scene_type, [0, 0])
        item[0] += 1
        decision = recommendations.get(_packet_key(packet))
        if decision and decision.recommended:
            item[1] += 1
    return {scene: (values[0], values[1]) for scene, values in counts.items()}


def _scene_card_sort_key(
    packet: SitePacket,
    decision: RecommendationDecision | None,
    *,
    image_assets: dict[str, Path] | None = None,
) -> tuple:
    return (
        (
            0
            if image_assets is None
            or _packet_has_image_asset(packet, image_assets)
            else 1
        ),
        0 if packet.entity.hero_image else 1,
        0 if decision and decision.recommended else 1,
        0 if decision and decision.metric_value is not None else 1,
        decision.rank if decision and decision.rank is not None else 10_000,
        -(packet.scene.annual_visits_est or 0),
        packet.entity.property_name.casefold(),
    )


def _packet_has_image_asset(
    packet: SitePacket,
    image_assets: dict[str, Path],
) -> bool:
    hero = packet.entity.hero_image
    return bool(hero and str(hero.url) in image_assets)


def _recommendation_mix(
    counts: dict[str, tuple[int, int]],
    locale: str,
) -> str:
    ranked = sorted(
        (
            (scene, recommended)
            for scene, (_candidate, recommended) in counts.items()
            if recommended > 0
        ),
        key=lambda item: (-item[1], _SCENE_ORDER.index(item[0]) if item[0] in _SCENE_ORDER else 99),
    )[:4]
    if not ranked:
        return "暂无门槛推荐" if locale == "zh" else "No threshold recommendations"
    return "  ·  ".join(
        f"{scene_card_label(scene, locale)} {count}"
        for scene, count in ranked
    )


def _card_metric_text(
    packet: SitePacket,
    decision: RecommendationDecision,
    locale: str,
) -> str:
    if decision.metric_value is None:
        return "硬主指标待补" if locale == "zh" else "Hard metric pending"
    label = metric_label(decision.metric_key, locale)
    return f"{label}: {_format_metric_value(decision.metric_value)}"


def _card_threshold_text(
    decision: RecommendationDecision,
    locale: str,
) -> str:
    prefix = "门槛：" if locale == "zh" else "Threshold: "
    text = decision.threshold_text
    if text.startswith("门槛：") or text.startswith("Threshold:"):
        return text
    return f"{prefix}{text}"


def _card_status(
    decision: RecommendationDecision,
    locale: str,
) -> tuple[str, RGBColor, RGBColor]:
    if decision.recommended:
        return (
            ("达标" if locale == "zh" else "QUALIFIED"),
            _PALE_GREEN,
            RGBColor(0, 132, 90),
        )
    if decision.metric_value is None:
        return (
            ("主指标缺失" if locale == "zh" else "METRIC MISSING"),
            _PALE_YELLOW,
            RGBColor(138, 104, 0),
        )
    return (
        ("未达标" if locale == "zh" else "BELOW THRESHOLD"),
        RGBColor(241, 243, 246),
        _GRAY,
    )


def _localized_country_name(
    country: str,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> str:
    if locale == "en" or not country:
        return country
    return localize_text(
        country,
        text_kind="country_name",
        target_locale=locale,
        cache=localization_cache,
        allow_provider=False,
        fallback_text=country,
    ).translated_text


def _title_font_size(title: str, locale: str) -> float:
    if locale == "zh":
        return 27 if len(title) >= 13 else 30
    return 24 if len(title) >= 40 else 28


def _property_font_size(name: str, locale: str) -> float:
    if len(name) >= 52:
        return 8.5
    if len(name) >= 38:
        return 9.5
    return 10.5 if locale == "en" else 11


def _card_metric_font_size(text: str, *, locale: str, dense: bool) -> float:
    if not dense:
        return 13.5
    if locale == "en" and len(text) >= 32:
        return 9.25
    if locale == "en" and len(text) >= 24:
        return 10.5
    return 12.5


def _set_shape_text(shape, value: str, *, size: float | None = None) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = value
    if size is not None:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def _set_shape_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _set_shape_font_color(shape, color: RGBColor) -> None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color


def _apply_brand_flag(shapes: dict[str, object]) -> None:
    _set_shape_fill(shapes["flag-red"], RGBColor(215, 25, 45))
    _set_shape_fill(shapes["flag-gold"], RGBColor(231, 185, 0))
    _set_shape_fill(shapes["flag-green"], RGBColor(0, 132, 90))
    _set_shape_text(shapes["flag-star"], "")


def _replace_picture(slide, picture, image_path: Path) -> None:
    _image_part, relationship_id = slide.part.get_or_add_image_part(str(image_path))
    picture._element.blipFill.blip.rEmbed = relationship_id


def _reflow_scene_cards(slide, representatives: list[SitePacket], pictures: list) -> None:
    count = len(representatives)
    if count == 0:
        return
    frames = _card_frames(count)
    for index, frame in enumerate(frames, start=1):
        _transform_card_group(slide, pictures[index - 1], index, frame)
    for index in range(count + 1, 10):
        suffix = f"{index:02d}"
        for shape in list(slide.shapes):
            if shape.name == f"card-{suffix}" or shape.name.endswith(f"-{suffix}"):
                _remove_shape(shape)
        _remove_shape(pictures[index - 1])


def _card_frames(count: int) -> list[tuple[float, float, float, float]]:
    left = 0.35
    width = 12.63
    gap = 0.12
    if count <= 5:
        card_width = (width - gap * (count - 1)) / count
        return [
            (left + index * (card_width + gap), 2.18, card_width, 3.72)
            for index in range(count)
        ]
    top_count = (count + 1) // 2
    bottom_count = count - top_count
    card_height = 2.5
    row_gap = 0.12
    frames: list[tuple[float, float, float, float]] = []
    for row, columns in enumerate((top_count, bottom_count)):
        card_width = (width - gap * (columns - 1)) / columns
        row_width = card_width * columns + gap * (columns - 1)
        row_left = left + (width - row_width) / 2
        for column in range(columns):
            frames.append(
                (
                    row_left + column * (card_width + gap),
                    1.92 + row * (card_height + row_gap),
                    card_width,
                    card_height,
                )
            )
    return frames


def _transform_card_group(
    slide,
    picture,
    index: int,
    frame: tuple[float, float, float, float],
) -> None:
    suffix = f"{index:02d}"
    card = next(shape for shape in slide.shapes if shape.name == f"card-{suffix}")
    old_left, old_top, old_width, old_height = (
        card.left,
        card.top,
        card.width,
        card.height,
    )
    new_left, new_top, new_width, new_height = (Inches(value) for value in frame)
    scale_x = new_width / old_width
    scale_y = new_height / old_height
    group = [
        shape
        for shape in slide.shapes
        if shape.name == f"card-{suffix}" or shape.name.endswith(f"-{suffix}")
    ]
    group.append(picture)
    for shape in group:
        relative_left = shape.left - old_left
        relative_top = shape.top - old_top
        shape.left = int(new_left + relative_left * scale_x)
        shape.top = int(new_top + relative_top * scale_y)
        shape.width = int(shape.width * scale_x)
        shape.height = int(shape.height * scale_y)


def _remove_shape(shape) -> None:
    element = shape._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _resolve_ppt_image(
    url: str,
    local_evidence_image: Path | None = None,
) -> Path | None:
    candidates = list(
        dict.fromkeys(
            [url, stable_hero_image_url(url), *hero_image_url_variants(url)]
        )
    )
    seen: set[str] = set()
    for candidate in candidates:
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        cached = _ppt_image_cache_path(candidate)
        if cached.exists() and cached.stat().st_size > 0:
            normalized = _normalize_ppt_image(cached)
            if normalized is not None:
                return normalized
    target = _ppt_image_cache_path(url)
    if local_evidence_image is not None and _looks_like_image(local_evidence_image):
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(local_evidence_image, target)
        target.with_suffix(".json").write_text(
            json.dumps(
                {
                    "content_type": "image/local-evidence-cache",
                    "final_url": url,
                    "local_source": str(local_evidence_image),
                }
            ),
            encoding="utf-8",
        )
        normalized = _normalize_ppt_image(target)
        if normalized is not None:
            return normalized
    headers = {
        "Accept": "image/png,image/jpeg",
        "User-Agent": "Mozilla/5.0 iSite2 standard report",
    }
    for candidate in candidates:
        try:
            with httpx.Client(
                follow_redirects=True,
                timeout=httpx.Timeout(6.0, connect=3.0),
                trust_env=False,
                headers=headers,
            ) as client:
                response = client.get(candidate)
            content_type = response.headers.get("content-type", "").casefold()
            if response.status_code not in {200, 206} or not content_type.startswith("image/"):
                continue
            if not response.content or len(response.content) > 12 * 1024 * 1024:
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(response.content)
            target.with_suffix(".json").write_text(
                json.dumps(
                    {
                        "content_type": content_type.split(";", 1)[0],
                        "final_url": str(response.url),
                    }
                ),
                encoding="utf-8",
            )
            normalized = _normalize_ppt_image(target)
            if normalized is not None:
                return normalized
        except (httpx.HTTPError, OSError):
            continue
    return None


def _ppt_image_cache_path(url: str) -> Path:
    cache_root = Path(
        os.getenv(
            "ISITE2_HERO_IMAGE_CACHE_DIR",
            str(_ROOT / ".tmp" / "isite2_hero_image_cache"),
        )
    )
    key = hashlib.sha256(url.encode("utf-8")).hexdigest()
    return cache_root / key[:2] / key[2:4] / f"{key}.image"


def _local_evidence_image_index() -> list[tuple[str, Path]]:
    root = _ROOT / ".web_evidence"
    if not root.exists():
        return []
    return [
        (_normalize_image_identity(path.stem), path)
        for path in root.glob("**/image_checks/*")
        if path.is_file() and path.stat().st_size > 0
    ]


def _matching_local_evidence_image(
    property_name: str,
    index: list[tuple[str, Path]],
) -> Path | None:
    identity = _normalize_image_identity(property_name)
    if not identity:
        return None
    matches = [
        path
        for file_identity, path in index
        if identity in file_identity or file_identity.endswith(identity)
    ]
    return max(matches, key=lambda path: path.stat().st_size) if matches else None


def _normalize_image_identity(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.casefold()).lstrip("0123456789")


def _looks_like_image(path: Path) -> bool:
    try:
        header = path.read_bytes()[:16]
    except OSError:
        return False
    return (
        header.startswith(b"\xff\xd8\xff")
        or header.startswith(b"\x89PNG\r\n\x1a\n")
        or header.startswith((b"GIF87a", b"GIF89a"))
        or header.startswith(b"RIFF") and header[8:12] == b"WEBP"
    )


def _normalize_ppt_image(path: Path) -> Path | None:
    """Return an image format accepted by python-pptx, converting when needed."""
    normalized = path.with_name(f"{path.stem}.ppt.jpg")
    try:
        with path.open("rb") as source:
            header = source.read(262144)
        image_format = _image_format_from_magic(header)
        if image_format in {"BMP", "GIF", "JPEG", "PNG", "TIFF"}:
            return path
        if image_format != "MPO":
            return None
        with Image.open(path) as image:
            if normalized.exists() and normalized.stat().st_size > 0:
                return normalized
            image.seek(0)
            frame = ImageOps.exif_transpose(image).convert("RGB")
            normalized.parent.mkdir(parents=True, exist_ok=True)
            frame.save(normalized, format="JPEG", quality=90, optimize=True)
            return normalized
    except (OSError, ValueError):
        return None


def _image_format_from_magic(content: bytes) -> str | None:
    if content.startswith(b"\xff\xd8\xff"):
        return "MPO" if b"MPF\x00" in content else "JPEG"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "PNG"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "GIF"
    if content.startswith(b"BM"):
        return "BMP"
    if content.startswith((b"II*\x00", b"MM\x00*")):
        return "TIFF"
    if content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "WEBP"
    if len(content) >= 12 and content[4:8] == b"ftyp":
        if content[8:12] in {b"avif", b"avis"}:
            return "AVIF"
    return None


def _add_methodology_slide(deck: Presentation, locale: str) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_title(
        slide,
        "AI工具——找楼逻辑" if locale == "zh" else "AI Tool - Building Discovery Logic",
        0.35,
    )
    _add_text(
        slide,
        (
            "基于公开证据、分场景硬主指标和质量门生成候选楼宇清单"
            if locale == "zh"
            else "Build the candidate site list from public evidence, scene-specific hard metrics, and quality gates"
        ),
        0.6,
        0.82,
        12.1,
        0.3,
        size=15,
        color=_GRAY,
    )

    steps = (
        [
            ("1", "从目录页找线索", "维基目录、官方清单和行业目录优先"),
            ("2", "补齐地理与硬证据", "核验经纬度、真实图片和场景主指标"),
            ("3", "Proxy模型估算价值", "推算访问量、忙时用户、流量和推荐理由"),
            ("4", "质量门形成清单", "通过证据、图片、重复和坐标检查后入表"),
        ]
        if locale == "zh"
        else [
            ("1", "Find Building Leads", "Start from wiki, official, and industry directories"),
            ("2", "Add Geo + Evidence", "Verify coordinates, real images, and scene metrics"),
            ("3", "Estimate Value by Proxy", "Derive visits, busy-hour users, traffic, and rationale"),
            ("4", "Pass Quality Gates", "Only qualified evidence, images, coordinates, and de-dup enter the list"),
        ]
    )
    for index, (number, header, body) in enumerate(steps):
        x = 0.6 + index * 3.15
        _add_process_box(slide, x, 1.2, number, header, body)
        if index < len(steps) - 1:
            _add_arrow(slide, x + 2.75, 1.58)

    for index, scene_type in enumerate(_METHODOLOGY_SCENES):
        row = index // 4
        col = index % 4
        _add_scene_rule_box(
            slide,
            scene_type,
            locale,
            0.55 + col * 3.17,
            2.55 + row * 2.08,
            2.85,
            1.78,
        )
    _add_text(
        slide,
        (
            "其余指标仅用于物业点价值评估；推荐清单严格优先使用场景级硬主指标。"
            if locale == "zh"
            else "Other indicators are used only for site value assessment; recommendations prioritize scene-level hard metrics."
        ),
        0.65,
        6.86,
        12.0,
        0.3,
        size=12,
        color=_GRAY,
    )


def _add_scene_statistics_slide(
    deck: Presentation,
    packets: list[SitePacket],
    recommendations: dict,
    locale: str,
) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    countries = _country_names(packets)
    country_title = _country_title(countries, locale)
    _add_title(
        slide,
        (
            f"{country_title}扫网结果与推荐点"
            if locale == "zh"
            else f"{country_title} Discovery Results & Recommended Sites"
        ),
        0.35,
    )
    recommended_packets = _recommended_packets(packets, recommendations)
    _add_summary_card(
        slide,
        0.6,
        0.92,
        "候选点" if locale == "zh" else "Candidates",
        len(packets),
        _LIGHT_BLUE,
    )
    _add_summary_card(
        slide,
        3.05,
        0.92,
        "推荐点" if locale == "zh" else "Recommended",
        len(recommended_packets),
        _PALE_YELLOW,
    )
    _add_summary_card(
        slide,
        5.5,
        0.92,
        "覆盖场景" if locale == "zh" else "Scenes Covered",
        len({packet.entity.scene_type for packet in packets}),
        _PALE_GREEN,
    )

    _add_text(
        slide,
        "分场景统计" if locale == "zh" else "Scene Statistics",
        0.6,
        1.78,
        5.6,
        0.35,
        size=18,
        bold=True,
        color=_BLUE,
    )
    _add_scene_stat_rows(slide, packets, recommendations, locale, 0.6, 2.18, 5.65, 4.78)

    _add_text(
        slide,
        "重点推荐物业点" if locale == "zh" else "Priority Recommended Sites",
        6.55,
        1.78,
        5.9,
        0.35,
        size=18,
        bold=True,
        color=_BLUE,
    )
    _add_recommendation_rows(
        slide,
        recommended_packets[:10],
        recommendations,
        locale,
        6.55,
        2.18,
        6.25,
        4.95,
    )


def _add_scene_stat_rows(
    slide,
    packets: list[SitePacket],
    recommendations: dict,
    locale: str,
    x: float,
    y: float,
    w: float,
    h: float,
) -> None:
    scene_counts: dict[str, int] = {}
    recommended_counts: dict[str, int] = {}
    for packet in packets:
        scene_counts[packet.entity.scene_type] = scene_counts.get(packet.entity.scene_type, 0) + 1
        decision = recommendations.get(_packet_key(packet))
        if decision and decision.recommended:
            recommended_counts[packet.entity.scene_type] = (
                recommended_counts.get(packet.entity.scene_type, 0) + 1
            )
    scenes = [scene for scene in _SCENE_ORDER if scene in scene_counts]
    scenes.extend(sorted(set(scene_counts) - set(scenes)))
    if not scenes:
        _add_text(slide, t("fallback.no_available_metric", locale), x, y, w, 0.4, size=12)
        return

    max_count = max(scene_counts.values()) or 1
    row_h = min(0.43, h / max(len(scenes), 1))
    for index, scene_type in enumerate(scenes):
        yy = y + index * row_h
        count = scene_counts.get(scene_type, 0)
        rec = recommended_counts.get(scene_type, 0)
        _add_text(
            slide,
            scene_label(scene_type, locale),
            x,
            yy + 0.02,
            1.55,
            row_h - 0.03,
            size=10.5,
            bold=True,
            color=_GRAY,
        )
        bar_w = max(0.18, (w - 2.85) * count / max_count)
        _add_rect(slide, x + 1.65, yy + 0.09, bar_w, row_h - 0.16, _LIGHT_BLUE, None)
        _add_text(
            slide,
            (
                f"候选 {count} / 推荐 {rec}"
                if locale == "zh"
                else f"{count} candidates / {rec} recommended"
            ),
            x + 1.75,
            yy + 0.01,
            w - 1.8,
            row_h - 0.02,
            size=10.5,
            color=RGBColor(46, 46, 46),
        )


def _add_recommendation_rows(
    slide,
    packets: list[SitePacket],
    recommendations: dict,
    locale: str,
    x: float,
    y: float,
    w: float,
    h: float,
) -> None:
    if not packets:
        _add_text(
            slide,
            "暂无达到推荐门槛的物业点" if locale == "zh" else "No sites meet the recommendation threshold.",
            x,
            y,
            w,
            0.5,
            size=12,
            color=_GRAY,
        )
        return
    row_h = h / min(10, len(packets))
    for index, packet in enumerate(packets):
        yy = y + index * row_h
        if index % 2 == 0:
            _add_rect(slide, x, yy, w, row_h - 0.03, RGBColor(248, 250, 252), _BORDER)
        decision = recommendations[_packet_key(packet)]
        rank = decision.rank if decision.rank is not None else index + 1
        metric = _compact_metric_text(packet, locale)
        _add_text(
            slide,
            f"{rank}. {packet.entity.property_name}",
            x + 0.08,
            yy + 0.05,
            3.12,
            row_h - 0.08,
            size=9.8,
            bold=True,
            color=RGBColor(32, 32, 32),
        )
        _add_text(
            slide,
            scene_label(packet.entity.scene_type, locale),
            x + 3.28,
            yy + 0.05,
            1.18,
            row_h - 0.08,
            size=9.2,
            color=_BLUE,
        )
        _add_text(
            slide,
            metric,
            x + 4.42,
            yy + 0.05,
            w - 4.5,
            row_h - 0.08,
            size=9.2,
            color=_GRAY,
        )


def _add_process_box(slide, x: float, y: float, number: str, header: str, body: str) -> None:
    _add_rect(slide, x, y, 2.52, 0.98, _LIGHT_BLUE, _BLUE)
    _add_text(slide, number, x + 0.12, y + 0.18, 0.32, 0.42, size=18, bold=True, color=_BLUE)
    _add_text(slide, header, x + 0.48, y + 0.12, 1.95, 0.28, size=12.2, bold=True, color=_BLUE)
    _add_text(slide, body, x + 0.48, y + 0.43, 1.92, 0.42, size=9.3, color=_GRAY)


def _add_scene_rule_box(
    slide,
    scene_type: str,
    locale: str,
    x: float,
    y: float,
    w: float,
    h: float,
) -> None:
    _add_rect(slide, x, y, w, h, RGBColor(255, 255, 255), _BORDER)
    _add_text(
        slide,
        scene_label(scene_type, locale),
        x + 0.12,
        y + 0.1,
        w - 0.24,
        0.25,
        size=13,
        bold=True,
        color=_BLUE,
    )
    _add_text(
        slide,
        _threshold_text(scene_type, locale),
        x + 0.12,
        y + 0.42,
        w - 0.24,
        0.32,
        size=11,
        bold=True,
        color=RGBColor(156, 101, 0),
    )
    _add_text(
        slide,
        "指标优先级" if locale == "zh" else "Metric priority",
        x + 0.12,
        y + 0.83,
        w - 0.24,
        0.22,
        size=10,
        bold=True,
        color=_GRAY,
    )
    indicators = _indicator_priority_text(scene_type, locale)
    _add_text(
        slide,
        indicators,
        x + 0.12,
        y + 1.08,
        w - 0.24,
        h - 1.16,
        size=8.7,
        color=RGBColor(60, 60, 60),
    )


def _add_summary_card(slide, x: float, y: float, label: str, value: int, fill: RGBColor) -> None:
    _add_rect(slide, x, y, 2.1, 0.66, fill, _BORDER)
    _add_text(slide, f"{value:,}", x + 0.12, y + 0.08, 0.95, 0.42, size=23, bold=True, color=_BLUE)
    _add_text(slide, label, x + 1.08, y + 0.19, 0.9, 0.28, size=12, bold=True, color=_GRAY)


def _add_title(slide, title: str, y: float) -> None:
    _add_text(slide, title, 0.6, y, 12.2, 0.42, size=28, bold=True, color=_BLUE)


def _add_arrow(slide, x: float, y: float) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(0.28), Inches(0.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _BLUE
    shape.line.color.rgb = _BLUE


def _add_rect(slide, x: float, y: float, w: float, h: float, fill: RGBColor, line: RGBColor | None) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.75)


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    bold: bool = False,
    color: RGBColor = RGBColor(0, 0, 0),
    align: PP_ALIGN | None = None,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def _country_names(packets: list[SitePacket]) -> list[str]:
    countries = []
    seen = set()
    for packet in packets:
        country = packet.entity.country
        if country not in seen:
            seen.add(country)
            countries.append(country)
    return countries


def _country_title(countries: list[str], locale: str) -> str:
    if not countries:
        return "国家" if locale == "zh" else "Country"
    if len(countries) == 1:
        return countries[0]
    if len(countries) <= 4:
        return "、".join(countries) if locale == "zh" else ", ".join(countries)
    return "多国" if locale == "zh" else "Multi-country"


def _threshold_text(scene_type: str, locale: str) -> str:
    gates = (
        load_output_template()
        .get("excel", {})
        .get("recommendation_rules", {})
        .get("default_scene_gates", {})
    )
    gate = gates.get(scene_type) or {}
    text = gate.get(f"threshold_text_{locale}") or gate.get("threshold_text_en")
    prefix = "First Class：" if locale == "zh" else "First Class: "
    return f"{prefix}{text}" if text else f"{prefix}{t('fallback.no_available_metric', locale)}"


def _indicator_priority_text(scene_type: str, locale: str) -> str:
    scene = scene_definitions().get(scene_type, {})
    indicators = list(scene.get("primary_indicators", []))[:3]
    if not indicators:
        return t("fallback.no_available_metric", locale)
    return "\n".join(
        f"{index}. {metric_label(indicator, locale)}"
        for index, indicator in enumerate(indicators, start=1)
    )


def _compact_metric_text(packet: SitePacket, locale: str) -> str:
    evidence = _select_main_evidence(packet)
    if evidence is None:
        return t("fallback.missing_primary_metric", locale)
    metric_key = _evidence_metric_key(evidence, packet.entity.scene_type)
    value = _format_metric_value(_metric_sort_value(evidence.field_value, metric_key))
    label = metric_label(evidence.indicator_name or evidence.field_group, locale)
    if value:
        return f"{label}: {value}"
    text = _main_metric_text(evidence, locale)
    return text[:96] if len(text) <= 96 else text[:96]


def _metric_value_for_display(value: str | None) -> float | None:
    if not value:
        return None
    matches = list(
        re.finditer(
            r"(?<![A-Za-z0-9])(\d+(?:[,.]\d{3})*(?:\.\d+)?|\d+(?:\.\d+)?)\s*(million|m|milhão|milhões)?",
            value,
            flags=re.IGNORECASE,
        )
    )
    if not matches:
        return None
    best = max(matches, key=lambda item: _display_number(item.group(1), item.group(2)))
    return _display_number(best.group(1), best.group(2))


def _display_number(number: str, multiplier: str | None) -> float:
    value = float(number.replace(",", ""))
    if multiplier and multiplier.lower() in {"million", "m", "milhão", "milhões"}:
        value *= 1_000_000
    return value


def _add_site_card(
    deck: Presentation,
    packet: SitePacket,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    entity = packet.entity
    demand = packet.demand
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9.0), Inches(0.5))
    title_frame = title_box.text_frame
    title_frame.text = entity.property_name
    title_frame.paragraphs[0].font.size = Pt(24)
    title_frame.paragraphs[0].font.bold = True

    core_metric_evidence = _select_main_evidence(packet)
    core_metric = (
        _main_metric_text(core_metric_evidence, locale)
        if core_metric_evidence is not None
        else t("fallback.unknown", locale)
    )
    busy_users = _fmt(demand.busy_hour_users if demand else None, locale)
    busy_traffic = _fmt(demand.busy_hour_traffic_gb if demand else None, locale)
    busy_bandwidth = _fmt(demand.busy_hour_bandwidth_mbps if demand else None, locale)
    coordinates = (
        f"{ppt_text('coordinates', locale)}: "
        f"{entity.latitude}, {entity.longitude} ({entity.geocode_precision})"
    )
    value_judgment = _localized_text(
        packet.conclusion.reason_to_recommend,
        "reason_to_recommend",
        locale,
        localization_cache,
    )
    risk_next_action = _localized_text(
        packet.conclusion.next_action,
        "next_action",
        locale,
        localization_cache,
    )

    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9.0), Inches(5.8))
    body.text_frame.text = "\n".join(
        [
            f"{ppt_text('location', locale)}: {entity.country} / {entity.city}",
            f"{ppt_text('scene', locale)}: {scene_label(entity.scene_type, locale)}",
            (
                f"{ppt_text('recommended_solution', locale)}: "
                f"{enum_label(packet.conclusion.recommended_solution, locale)}"
            ),
            (
                f"{ppt_text('status', locale)}: "
                f"{enum_label(packet.conclusion.evidence_status, locale)} / "
                f"{enum_label(packet.conclusion.value_class, locale)} / "
                f"{enum_label(packet.conclusion.action_class, locale)}"
            ),
            (f"{ppt_text('core_metric', locale)}: {core_metric}"),
            (
                f"{ppt_text('busy_hour', locale)}: "
                f"users={busy_users}, traffic_gb={busy_traffic}, "
                f"bandwidth_mbps={busy_bandwidth}"
            ),
            (
                f"{ppt_text('area_scale', locale)}: "
                f"{packet.scene.area_metric_name} ({packet.scene.area_metric_status})"
            ),
            coordinates,
            (
                f"{ppt_text('build_status', locale)}: "
                f"{enum_label(packet.build_status.indoor_system_presence, locale)} / "
                f"{enum_label(packet.build_status.indoor_system_type, locale)} / "
                f"{enum_label(packet.build_status.indoor_rat, locale)}"
            ),
            f"{ppt_text('value_judgment', locale)}: {value_judgment}",
            f"{ppt_text('risk_next_action', locale)}: {risk_next_action}",
            f"{ppt_text('evidence_entries', locale)}: {len(packet.evidence)}",
        ]
    )


def _fmt(value: float | None, locale: str) -> str:
    if value is None:
        return t("fallback.unknown", locale)
    return f"{value:.2f}"


def _localized_text(
    source_text: str,
    text_kind: str,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> str:
    return localize_text(
        source_text,
        text_kind=text_kind,
        target_locale=locale,
        cache=localization_cache,
        allow_provider=False,
        fallback_text=generic_free_text_fallback(text_kind, locale),
    ).translated_text
