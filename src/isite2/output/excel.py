from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from isite2.domain.models import EvidenceItem, SitePacket
from isite2.localization import (
    DatabaseLocalizationCache,
    enum_label,
    excel_labels,
    generic_free_text_fallback,
    localize_text,
    metric_label,
    resolve_locale,
    scene_label,
    t,
)
from isite2.rules.candidate_quality import filter_packets_for_surface
from isite2.rules.config_loader import (
    load_candidate_quality_rules,
    load_output_template,
    scene_definitions,
)
from isite2.rules.metric_identity import (
    MetricIdentityProvider,
    accepted_metric_identity,
    metric_identity_request,
    recommendation_metric_identity_provider_from_env,
)

MAIN_SHEET_KEY = "main"
SCENE_SHEET_KEYS = {
    "airport_terminal",
    "convention_center",
    "mall_mixed_use",
    "stadium",
    "luxury_hotel_mice",
    "office_government",
    "hospital",
    "university",
    "transport_hub",
    "mosque",
}
RECOMMENDATION_TARGET_COUNT = 10
RECOMMENDATION_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
RECOMMENDATION_FONT = Font(color="7A4F01")


@dataclass(frozen=True)
class RecommendationDecision:
    recommended: bool
    rank: int | None
    metric_key: str
    metric_value: float | None
    threshold_metric_key: str
    threshold_value: float | None
    threshold_text: str
    hard_metric_candidate_count: int
    recommended_count: int


@dataclass(frozen=True)
class _RecommendationScore:
    packet: SitePacket
    metric_key: str
    metric_value: float
    metric_order: int
    source_tier_rank: int
    source_date_rank: int


def create_excel_workbook(
    packets: list[SitePacket] | None = None,
    scope_text: str = "MVP skeleton",
    include_blocked_quality: bool = False,
    metric_identity_provider: MetricIdentityProvider | None = None,
    locale: str | None = None,
    localization_cache: DatabaseLocalizationCache | None = None,
    recommendations: dict[str, RecommendationDecision] | None = None,
) -> Workbook:
    resolved_locale = resolve_locale(locale)
    packets = _sort_main_packets(
        filter_packets_for_surface(
            packets or [],
            "export",
            include_blocked_quality=include_blocked_quality,
        )
    )
    excel = excel_labels(resolved_locale)
    metric_identity_provider = (
        metric_identity_provider
        if metric_identity_provider is not None
        else recommendation_metric_identity_provider_from_env()
    )
    recommendations = recommendations or _build_recommendation_index(
        packets,
        metric_identity_provider=metric_identity_provider,
        locale=resolved_locale,
    )

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for sheet_key in excel["required_sheet_keys"]:
        sheet_name = excel["sheets"][sheet_key]
        worksheet = workbook.create_sheet(sheet_name)
        _write_sheet(
            worksheet,
            sheet_key,
            excel,
            packets,
            scope_text,
            recommendations,
            locale=resolved_locale,
            localization_cache=localization_cache,
        )
        _format_sheet_layout(worksheet)
    return workbook


def _format_sheet_layout(worksheet) -> None:
    def display_width(text: str) -> int:
        return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in text)

    for column in worksheet.iter_cols():
        longest = max(
            (display_width(line) for cell in column for line in str(cell.value or "").splitlines()),
            default=0,
        )
        worksheet.column_dimensions[column[0].column_letter].width = min(60, max(14, longest + 3))
    for row in worksheet.iter_rows():
        line_count = 1
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            width = worksheet.column_dimensions[cell.column_letter].width - 3
            # Leave room for word wrapping and wide-script glyphs without changing values.
            lines = sum(
                max(1, math.ceil(display_width(line) * 1.2 / width))
                for line in str(cell.value or "").splitlines()
            )
            line_count = max(line_count, lines)
        worksheet.row_dimensions[row[0].row].height = min(409, 15 * line_count + 6)
    worksheet.freeze_panes = "A2"


def write_excel_skeleton(
    path: Path,
    packets: list[SitePacket] | None = None,
    include_blocked_quality: bool = False,
    metric_identity_provider: MetricIdentityProvider | None = None,
    locale: str | None = None,
    localization_cache: DatabaseLocalizationCache | None = None,
    recommendations: dict[str, RecommendationDecision] | None = None,
) -> Path:
    workbook = create_excel_workbook(
        packets=packets,
        include_blocked_quality=include_blocked_quality,
        metric_identity_provider=metric_identity_provider,
        locale=locale,
        localization_cache=localization_cache,
        recommendations=recommendations,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _write_sheet(
    worksheet,
    sheet_key: str,
    excel: dict,
    packets: list[SitePacket],
    scope_text: str,
    recommendations: dict[str, RecommendationDecision],
    *,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> None:
    if sheet_key == MAIN_SHEET_KEY or sheet_key in SCENE_SHEET_KEYS:
        worksheet.append(excel["columns"]["main"])
        for packet in _packets_for_sheet(sheet_key, packets):
            worksheet.append(
                _main_row(
                    packet,
                    recommendations.get(_packet_key(packet)),
                    locale=locale,
                    localization_cache=localization_cache,
                )
            )
            _highlight_if_recommended(worksheet, worksheet.max_row, packet, recommendations)
    elif sheet_key == "recommendation":
        worksheet.append(excel["columns"]["recommendation"])
        for packet in _recommended_packets(packets, recommendations):
            decision = recommendations[_packet_key(packet)]
            worksheet.append(
                _recommendation_row(
                    packet,
                    decision,
                    locale=locale,
                    localization_cache=localization_cache,
                )
            )
            _highlight_row(worksheet, worksheet.max_row)
    elif sheet_key == "evidence":
        worksheet.append(excel["columns"]["evidence"])
        for row_number, packet in enumerate(packets, start=1):
            for evidence in packet.evidence:
                original = evidence.field_value
                output_original = _clean_locale_text(
                    original,
                    locale,
                    "evidence.field_value",
                )
                translation = _translate_for_output(
                    original,
                    "evidence.field_value",
                    locale,
                    localization_cache,
                )
                worksheet.append(
                    [
                        row_number,
                        packet.entity.country,
                        packet.entity.city,
                        packet.entity.property_name,
                        scene_label(packet.entity.scene_type, locale),
                        metric_label(evidence.field_group, locale),
                        output_original,
                        translation,
                        enum_label(evidence.evidence_type, locale),
                        enum_label(evidence.source_tier, locale),
                        str(evidence.source_url),
                        "",
                        enum_label(evidence.cross_check_status, locale),
                    ]
                )
    elif sheet_key == "inference":
        worksheet.append(excel["columns"]["inference"])
        for row_number, packet in enumerate(packets, start=1):
            for inference in packet.inference:
                worksheet.append(
                    [
                        row_number,
                        packet.entity.country,
                        packet.entity.property_name,
                        metric_label(inference.inferred_field, locale),
                        _translate_for_output(
                            inference.inferred_value,
                            "inference.inferred_value",
                            locale,
                            localization_cache,
                        ),
                        _translate_for_output(
                            inference.inference_basis,
                            "inference.inference_basis",
                            locale,
                            localization_cache,
                        ),
                        _translate_for_output(
                            inference.inference_chain,
                            "inference.inference_chain",
                            locale,
                            localization_cache,
                        ),
                        enum_label(inference.inference_confidence, locale),
                    ]
                )
    elif sheet_key == "proxy_model":
        worksheet.append(excel["columns"]["proxy_model"])
        for scene_type, scene in scene_definitions().items():
            params = scene["demand_parameters"]
            worksheet.append(
                [
                    scene_label(scene_type, locale),
                    " / ".join(metric_label(item, locale) for item in scene["primary_indicators"]),
                    metric_label(scene["area_metric"], locale),
                    " / ".join(scene["proxy_basis"]),
                    _range_text(params["attach_rate"]),
                    _range_text(params["indoor_capture"]),
                    _range_text(params["busy_hour_factor"]),
                    _range_text(params["gb_per_user_busy_hour"]),
                ]
            )
    elif sheet_key == "review_queue":
        worksheet.append(excel["columns"]["review_queue"])
        for row_number, packet in enumerate(packets, start=1):
            for review in packet.review_queue:
                worksheet.append(
                    [
                        row_number,
                        packet.entity.country,
                        packet.entity.city,
                        packet.entity.property_name,
                        _translate_for_output(
                            review.reason,
                            "review.reason",
                            locale,
                            localization_cache,
                        ),
                        _translate_for_output(
                            review.next_action,
                            "review.next_action",
                            locale,
                            localization_cache,
                        ),
                    ]
                )
    elif sheet_key == "country_summary":
        worksheet.append(excel["columns"]["country_summary"])
        _write_country_summary(worksheet, packets, locale)
    else:
        worksheet.append(["Item", "Description"] if locale == "en" else ["条目", "说明"])
        for row in _method_rows(scope_text, locale):
            worksheet.append(row)


def _packets_for_sheet(sheet_key: str, packets: list[SitePacket]) -> list[SitePacket]:
    if sheet_key == MAIN_SHEET_KEY:
        return packets
    scene_packets = [packet for packet in packets if packet.entity.scene_type == sheet_key]
    return _sort_scene_packets(scene_packets)


def _recommended_packets(
    packets: list[SitePacket],
    recommendations: dict[str, RecommendationDecision],
) -> list[SitePacket]:
    return sorted(
        [
            packet
            for packet in packets
            if recommendations.get(_packet_key(packet))
            and recommendations[_packet_key(packet)].recommended
        ],
        key=lambda packet: (
            str(packet.entity.country or "").casefold(),
            str(packet.entity.scene_type or "").casefold(),
            recommendations[_packet_key(packet)].rank or 10_000,
            str(packet.entity.city or "").casefold(),
            str(packet.entity.property_name or "").casefold(),
        ),
    )


def _main_row(
    packet: SitePacket,
    recommendation: RecommendationDecision | None = None,
    *,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> list:
    entity = packet.entity
    demand = packet.demand
    build = packet.build_status
    main_evidence = _select_main_evidence(packet)
    main_metric = _main_metric_text(main_evidence, locale)
    main_metric_value = _main_metric_numeric_value(packet, main_evidence)
    return [
        entity.country,
        entity.city,
        entity.property_name,
        entity.longitude,
        entity.latitude,
        scene_label(entity.scene_type, locale),
        main_metric,
        main_metric_value,
        packet.scene.annual_visits_est,
        demand.busy_hour_traffic_gb if demand else None,
        enum_label(build.indoor_system_presence, locale),
        enum_label(build.indoor_rat, locale),
        build.operator_name or t("fallback.operator_candidates", locale),
        enum_label(build.indoor_system_type, locale),
        "MVP rule-based batch",
        enum_label(packet.conclusion.evidence_status, locale),
        t("fallback.yes", locale) if recommendation and recommendation.recommended else "",
        _translate_for_output(
            packet.conclusion.reason_to_recommend,
            "reason_to_recommend",
            locale,
            localization_cache,
        ),
        _translate_for_output(
            packet.conclusion.next_action,
            "next_action",
            locale,
            localization_cache,
        ),
        entity.google_maps_link,
    ]


def _recommendation_row(
    packet: SitePacket,
    decision: RecommendationDecision,
    *,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> list:
    entity = packet.entity
    main_evidence = _select_main_evidence(packet)
    return [
        entity.country,
        entity.city,
        entity.property_name,
        scene_label(entity.scene_type, locale),
        decision.rank,
        decision.threshold_text,
        decision.metric_key,
        _format_metric_value(decision.metric_value),
        _main_metric_text(main_evidence, locale),
        packet.scene.annual_visits_est,
        _translate_for_output(
            packet.conclusion.reason_to_recommend,
            "reason_to_recommend",
            locale,
            localization_cache,
        ),
        entity.google_maps_link,
    ]


def _build_recommendation_index(
    packets: list[SitePacket],
    target_count: int = RECOMMENDATION_TARGET_COUNT,
    metric_identity_provider: MetricIdentityProvider | None = None,
    locale: str | None = None,
) -> dict[str, RecommendationDecision]:
    resolved_locale = resolve_locale(locale)
    recommendation_rules = load_output_template()["excel"].get("recommendation_rules", {})
    groups: dict[tuple[str, str], list[SitePacket]] = {}
    for packet in packets:
        groups.setdefault((packet.entity.country, packet.entity.scene_type), []).append(packet)

    decisions: dict[str, RecommendationDecision] = {}
    for (_country, scene_type), group_packets in groups.items():
        fixed_gate = _fixed_recommendation_gate(scene_type, recommendation_rules)
        if fixed_gate:
            scores = _recommendation_scores(
                group_packets,
                scene_type,
                preferred_metric_keys=set(fixed_gate["metric_keys"]),
                metric_identity_provider=metric_identity_provider,
            )
            threshold_metric_key = fixed_gate["metric_keys"][0]
            threshold_value = float(fixed_gate["value"])
            threshold_text = _format_fixed_threshold_text(fixed_gate, resolved_locale)
        else:
            scores = _recommendation_scores(
                group_packets,
                scene_type,
                metric_identity_provider=metric_identity_provider,
            )
            threshold_score = scores[min(target_count, len(scores)) - 1] if scores else None
            threshold_metric_key = threshold_score.metric_key if threshold_score else ""
            threshold_value = threshold_score.metric_value if threshold_score else None
            threshold_text = _format_threshold_text(
                threshold_metric_key,
                threshold_value,
                resolved_locale,
            )
        recommended_keys: set[str] = set()
        rank_by_key: dict[str, int] = {}
        score_by_key: dict[str, _RecommendationScore] = {}

        for rank, score in enumerate(scores, start=1):
            key = _packet_key(score.packet)
            rank_by_key[key] = rank
            score_by_key[key] = score
            if fixed_gate and _passes_fixed_gate(score.metric_value, fixed_gate):
                recommended_keys.add(key)
            elif (
                not fixed_gate
                and threshold_value is not None
                and (
                    rank <= target_count
                    or (
                        score.metric_key == threshold_metric_key
                        and score.metric_value >= threshold_value
                    )
                )
            ):
                recommended_keys.add(key)

        for packet in group_packets:
            key = _packet_key(packet)
            score = score_by_key.get(key)
            decisions[key] = RecommendationDecision(
                recommended=key in recommended_keys,
                rank=rank_by_key.get(key),
                metric_key=score.metric_key if score else threshold_metric_key,
                metric_value=score.metric_value if score else None,
                threshold_metric_key=threshold_metric_key,
                threshold_value=threshold_value,
                threshold_text=threshold_text,
                hard_metric_candidate_count=len(scores),
                recommended_count=len(recommended_keys),
            )
    return decisions


def _recommendation_scores(
    packets: list[SitePacket],
    scene_type: str,
    preferred_metric_keys: set[str] | None = None,
    metric_identity_provider: MetricIdentityProvider | None = None,
) -> list[_RecommendationScore]:
    metric_order = _scene_metric_order(scene_type)
    preferred = {
        _canonical_metric_key(_normalize_metric_key(metric))
        for metric in (preferred_metric_keys or set())
    }
    scores: list[_RecommendationScore] = []
    for packet in packets:
        score = _recommendation_score_for_packet(
            packet,
            scene_type,
            metric_order,
            preferred,
            metric_identity_provider,
        )
        if score is None:
            continue
        scores.append(score)
    return sorted(
        scores,
        key=lambda score: (
            score.metric_order,
            -score.metric_value,
            str(score.packet.entity.city or "").casefold(),
            str(score.packet.entity.property_name or "").casefold(),
        ),
    )


def _recommendation_score_for_packet(
    packet: SitePacket,
    scene_type: str,
    metric_order: dict[str, int],
    preferred_metric_keys: set[str],
    metric_identity_provider: MetricIdentityProvider | None = None,
) -> _RecommendationScore | None:
    evidence_items = packet.evidence
    if not preferred_metric_keys:
        main_evidence = _select_main_evidence(packet)
        evidence_items = [main_evidence] if main_evidence is not None else []

    candidates: list[_RecommendationScore] = []
    for evidence in evidence_items:
        if evidence is None or _has_blocked_objective_text(evidence):
            continue
        metric_key = _recommendation_metric_key(
            evidence,
            scene_type,
            metric_order,
            preferred_metric_keys,
            metric_identity_provider,
        )
        if metric_key not in metric_order and metric_key not in preferred_metric_keys:
            continue
        if _is_descriptive_metric(metric_key):
            continue
        if preferred_metric_keys and metric_key not in preferred_metric_keys:
            continue
        metric_value = _metric_sort_value(evidence.field_value, metric_key)
        if metric_value is None:
            continue
        if not _recommendation_metric_value_is_plausible(
            scene_type,
            metric_key,
            metric_value,
        ):
            continue
        candidates.append(
            _RecommendationScore(
                packet=packet,
                metric_key=metric_key,
                metric_value=metric_value,
                metric_order=metric_order.get(metric_key, 10_000),
                source_tier_rank=_source_tier_rank(str(evidence.source_tier)),
                source_date_rank=_source_date_rank(evidence.source_date),
            )
        )
    if not candidates:
        return None
    return sorted(
        candidates,
        key=lambda score: (
            score.metric_order,
            score.source_tier_rank,
            -score.source_date_rank,
            -score.metric_value,
            str(score.packet.entity.city or "").casefold(),
            str(score.packet.entity.property_name or "").casefold(),
        ),
    )[0]


def _recommendation_metric_value_is_plausible(
    scene_type: str,
    metric_key: str,
    metric_value: float,
) -> bool:
    normalized_key = _canonical_metric_key(_normalize_metric_key(metric_key))
    if (
        scene_type == "luxury_hotel_mice"
        and normalized_key in {"keys", "room_count"}
        and metric_value > 5_000
    ):
        return False
    return True


def _packet_key(packet: SitePacket) -> str:
    return str(packet.entity.property_id)


def _recommendation_metric_key(
    evidence: EvidenceItem,
    scene_type: str,
    metric_order: dict[str, int],
    preferred_metric_keys: set[str],
    metric_identity_provider: MetricIdentityProvider | None,
) -> str:
    rule_metric_key = _evidence_metric_key(evidence, scene_type, metric_order)
    field_value = str(evidence.field_value or "")
    if (
        scene_type == "office_government"
        and "tower_height" in preferred_metric_keys
        and _metric_value_for_key(field_value, "tower_height") is not None
    ):
        return "tower_height"
    if metric_identity_provider is None:
        return rule_metric_key
    allowed_metric_keys = {
        _canonical_metric_key(_normalize_metric_key(metric_key))
        for metric_key in {*metric_order.keys(), *preferred_metric_keys}
        if _normalize_metric_key(metric_key)
    }
    request = metric_identity_request(
        scene_type=scene_type,
        field_group=str(evidence.field_group or ""),
        indicator_name=str(evidence.indicator_name or ""),
        field_value=field_value,
        allowed_metric_keys=allowed_metric_keys,
        preferred_metric_keys=preferred_metric_keys,
    )
    request["rule_based_metric_key"] = rule_metric_key
    try:
        decision = metric_identity_provider.identify(request)
    except Exception:
        return rule_metric_key
    gpt_metric_key = accepted_metric_identity(
        decision,
        allowed_metric_keys=allowed_metric_keys,
    )
    return gpt_metric_key or rule_metric_key


def _highlight_if_recommended(
    worksheet,
    row_number: int,
    packet: SitePacket,
    recommendations: dict[str, RecommendationDecision],
) -> None:
    decision = recommendations.get(_packet_key(packet))
    if decision and decision.recommended:
        _highlight_row(worksheet, row_number)


def _highlight_row(worksheet, row_number: int) -> None:
    for cell in worksheet[row_number]:
        cell.fill = RECOMMENDATION_FILL
        cell.font = RECOMMENDATION_FONT


def _format_threshold_text(metric_key: str, value: float | None, locale: str | None = None) -> str:
    resolved_locale = resolve_locale(locale)
    if value is None or not metric_key:
        return t("fallback.no_available_metric", resolved_locale)
    return f"{metric_label(metric_key, resolved_locale)} >= {_format_metric_value(value)}"


def _fixed_recommendation_gate(scene_type: str, rules: dict) -> dict | None:
    gates = rules.get("default_scene_gates", {})
    gate = gates.get(scene_type)
    if not isinstance(gate, dict):
        return None
    metric_keys = [
        _canonical_metric_key(_normalize_metric_key(metric))
        for metric in gate.get("metric_keys", [])
        if _normalize_metric_key(metric)
    ]
    if not metric_keys or gate.get("value") is None:
        return None
    return {
        **gate,
        "metric_keys": metric_keys,
        "operator": str(gate.get("operator") or ">"),
        "value": float(gate["value"]),
    }


def _passes_fixed_gate(metric_value: float | None, gate: dict) -> bool:
    if metric_value is None:
        return False
    threshold = float(gate["value"])
    operator = str(gate.get("operator") or ">")
    if operator == ">":
        return metric_value > threshold
    if operator == ">=":
        return metric_value >= threshold
    return metric_value > threshold


def _format_fixed_threshold_text(gate: dict, locale: str | None = None) -> str:
    resolved_locale = resolve_locale(locale)
    localized_key = f"threshold_text_{resolved_locale}"
    return str(
        gate.get(localized_key)
        or gate.get("threshold_text_en")
        or _format_threshold_text(gate["metric_keys"][0], gate["value"], resolved_locale)
    )


def _format_metric_value(value: float | None) -> str:
    if value is None:
        return ""
    rounded = int(round(value))
    return f"{rounded:,}"


def _select_main_evidence(packet: SitePacket) -> EvidenceItem | None:
    """Return the strongest scene objective evidence for the Excel main table."""
    if not packet.evidence:
        return None
    scene_type = packet.entity.scene_type
    ranked = sorted(
        packet.evidence,
        key=lambda evidence: _evidence_priority(evidence, scene_type),
    )
    return ranked[0]


def _main_metric_text(evidence: EvidenceItem | None, locale: str | None = None) -> str:
    resolved_locale = resolve_locale(locale)
    if evidence is None:
        return t("fallback.missing_primary_metric", resolved_locale)
    metric_key = evidence.indicator_name or evidence.field_group
    return _clean_locale_text(
        f"{metric_label(metric_key, resolved_locale)}: {evidence.field_value}",
        resolved_locale,
        "evidence.field_value",
    )


def _main_metric_numeric_value(
    packet: SitePacket,
    evidence: EvidenceItem | None,
) -> int | float | None:
    if evidence is None:
        return None
    scene_type = packet.entity.scene_type
    if not _is_scene_objective_evidence(evidence, scene_type):
        return None
    metric_key = _evidence_metric_key(evidence, scene_type)
    value = _metric_sort_value(evidence.field_value, metric_key)
    if value is None:
        return None
    if float(value).is_integer():
        return int(value)
    return value


def _sort_main_packets(packets: list[SitePacket]) -> list[SitePacket]:
    return sorted(packets, key=_main_sort_key)


def _main_sort_key(packet: SitePacket) -> tuple[int, float, str, str, str]:
    annual_visits = _safe_float(packet.scene.annual_visits_est)
    return (
        1 if annual_visits is None else 0,
        -(annual_visits or 0.0),
        str(packet.entity.country or "").casefold(),
        str(packet.entity.city or "").casefold(),
        str(packet.entity.property_name or "").casefold(),
    )


def _sort_scene_packets(packets: list[SitePacket]) -> list[SitePacket]:
    return sorted(packets, key=_scene_sort_key)


def _scene_sort_key(packet: SitePacket) -> tuple[int, float, int, str, str, str]:
    evidence = _select_main_evidence(packet)
    metric_key = (
        _evidence_metric_key(evidence, packet.entity.scene_type) if evidence is not None else ""
    )
    metric_value = _metric_sort_value(evidence.field_value if evidence else "", metric_key)
    metric_order = _scene_metric_order(packet.entity.scene_type)
    return (
        1 if metric_value is None else 0,
        metric_order.get(metric_key, 10_000),
        -(metric_value or 0.0),
        str(packet.entity.country or "").casefold(),
        str(packet.entity.city or "").casefold(),
        str(packet.entity.property_name or "").casefold(),
    )


def _evidence_priority(evidence: EvidenceItem, scene_type: str) -> tuple[int, int, int, int, int]:
    metric_order = _scene_metric_order(scene_type)
    metric_key = _evidence_metric_key(evidence, scene_type, metric_order)
    accepted = metric_key in metric_order
    descriptive = _is_descriptive_metric(metric_key)
    blocked = _has_blocked_objective_text(evidence)
    has_number = bool(re.search(r"\d", evidence.field_value or ""))

    if accepted and has_number and not descriptive and not blocked:
        bucket = 0
    elif accepted and not descriptive and not blocked:
        bucket = 1
    elif accepted and has_number and not blocked:
        bucket = 2
    elif accepted and not blocked:
        bucket = 3
    elif has_number:
        bucket = 4
    else:
        bucket = 5

    return (
        bucket,
        metric_order.get(metric_key, 10_000),
        _evidence_type_rank(str(evidence.evidence_type)),
        _source_tier_rank(str(evidence.source_tier)),
        _cross_check_rank(str(evidence.cross_check_status)),
    )


def _is_scene_objective_evidence(evidence: EvidenceItem, scene_type: str) -> bool:
    metric_order = _scene_metric_order(scene_type)
    metric_key = _evidence_metric_key(evidence, scene_type, metric_order)
    return (
        metric_key in metric_order
        and not _is_descriptive_metric(metric_key)
        and not _has_blocked_objective_text(evidence)
    )


def _scene_metric_order(scene_type: str) -> dict[str, int]:
    rules = load_candidate_quality_rules()
    scene_rule = scene_definitions().get(scene_type, {})
    ordered_metrics = [
        *rules.get("objective_indicator_aliases", {}).get(scene_type, []),
        *scene_rule.get("primary_indicators", []),
        *scene_rule.get("fallback_metrics", []),
    ]
    metric_order: dict[str, int] = {}
    for metric in ordered_metrics:
        normalized = _normalize_metric_key(metric)
        if normalized and normalized not in metric_order:
            metric_order[normalized] = len(metric_order)
    return metric_order


def _evidence_metric_key(
    evidence: EvidenceItem,
    scene_type: str,
    metric_order: dict[str, int] | None = None,
) -> str:
    accepted = metric_order if metric_order is not None else _scene_metric_order(scene_type)
    keys = [
        _normalize_metric_key(evidence.indicator_name or ""),
        _normalize_metric_key(evidence.field_group),
    ]
    accepted_candidates = [
        canonical
        for key in [*keys, *_metric_keys_from_field_value(evidence.field_value)]
        if (canonical := _canonical_metric_key(key)) in accepted
    ]
    if accepted_candidates:
        return min(accepted_candidates, key=lambda candidate: accepted[candidate])
    for key in keys:
        if key:
            return _canonical_metric_key(key)
    return ""


def _metric_keys_from_field_value(value: str | None) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    prefixes = re.findall(r"(?:^|[;。])\s*([A-Za-z][A-Za-z0-9_ /-]{1,80})\s*:", text)
    keys: list[str] = []
    for prefix in prefixes:
        keys.append(_normalize_metric_key(prefix))
        keys.extend(
            _normalize_metric_key(part)
            for part in re.split(r"[/|]", prefix)
            if _normalize_metric_key(part)
        )
    return [key for key in keys if key]


def _canonical_metric_key(metric_key: str) -> str:
    aliases = {
        "abl": "gla",
        "rooms": "room_count",
        "guest_rooms": "room_count",
        "hotel_keys": "keys",
        "annual_visits": "annual_footfall",
        "annual_passengers": "annual_passenger_throughput",
        "passengers": "passenger_throughput",
        "gross_leasable_area": "gla",
        "gross_lettable_area": "gla",
        "gross_leaseable_area": "gla",
        "leasable_area": "gla",
        "lettable_area": "gla",
        "retail_gfa": "gla",
        "retail_floor_area": "gla",
        "retail_gross_floor_area": "gla",
        "student_count": "enrollment",
        "students": "enrollment",
        "visitor_count": "annual_visitors",
        "visitors": "annual_visitors",
        "annual_visitor_count": "annual_visitors",
        "built_up_area": "gross_floor_area",
        "building_area": "gross_floor_area",
    }
    return aliases.get(metric_key, metric_key)


def _normalize_metric_key(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", str(value).lower()).strip("_")
    return re.sub(r"_+", "_", normalized)


def _safe_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _metric_sort_value(text: str | None, metric_key: str | None = None) -> float | None:
    normalized_key = _canonical_metric_key(_normalize_metric_key(metric_key or ""))
    keyed_value = _metric_value_for_key(str(text or ""), normalized_key)
    if normalized_key in {
        "gla",
        "retail_gfa",
        "retail_floor_area",
        "exhibition_area",
        "meeting_area",
        "annual_passenger_throughput",
        "passenger_throughput",
        "beds",
        "keys",
        "room_count",
        "line_count",
        "tower_height",
        "seat_count",
        "peak_event_capacity",
        "plenary_capacity",
        "ballroom_capacity",
        "mosque_area",
        "gross_floor_area",
        "prayer_hall_area",
        "site_area",
        "built_up_area",
        "annual_visitors",
        "annual_visits",
        "daily_visitors",
        "enrollment",
    }:
        return keyed_value
    if keyed_value is not None:
        return keyed_value
    values = [_scaled_number(match) for match in _metric_number_pattern().finditer(str(text or ""))]
    values = [value for value in values if value is not None]
    if not values:
        return None
    return max(values)


def _metric_value_for_key(text: str, metric_key: str | None) -> float | None:
    key = _canonical_metric_key(_normalize_metric_key(metric_key or ""))
    lower = text.casefold()
    if key == "beds":
        return _first_context_number(
            lower,
            [
                r"\bbeds?\s*[:=]?\s*(\d[\d,.]*)",
                r"\b(\d[\d,.]*)\s*\+?\s*(?:[-\u2013\u2014]\s*)?beds?\b",
                r"\b(\d[\d,.]*)\s*床位\b",
            ],
        )
    if key in {"keys", "room_count"}:
        return _first_context_number(
            lower,
            [
                r"\b(?:guest\s+rooms?|rooms|keys?)\s*[:=]?\s*(\d[\d,.]*)",
                r"\b(\d[\d,.]*)\s*(?:guest\s*)?(?:rooms?|keys?)\b",
                r"\b(?:客房|房间数)\s*[:=]?\s*(\d[\d,.]*)",
            ],
            blocked_contexts=(
                "function room",
                "function rooms",
                "meeting room",
                "meeting rooms",
                "conference room",
                "conference rooms",
                "ballroom",
            ),
        )
    if key == "line_count":
        return _first_context_number(
            lower,
            [
                r"\bline\s*count\s*[:=]?\s*(\d{1,3})\b",
                r"\b(\d{1,3})\s*(?:metro\s*)?(?:rail\s*)?(?:lines|routes)\b",
                r"\b(\d{1,3})\s*connections?\b",
            ],
        )
    if key == "tower_height":
        return _first_context_number(
            lower,
            [
                r"\b(?:tower\s*)?(?:building\s*)?height\s*(?:\([^)]*m[^)]*\))?\s*[:=]?\s*(\d[\d,.]*)\s*(?:m|meters?|metres?)\b",
                r"\b(?:building\s*)?height\s*[:=]?\s*(\d[\d,.]*)\s*(?:m|meters?|metres?)\b",
                r"\b(\d[\d,.]*)\s*(?:m|meters?|metres?)\s*"
                r"(?:tall|height|tower height|building height)\b",
            ],
        )
    if key in {"seat_count", "peak_event_capacity", "plenary_capacity", "ballroom_capacity"}:
        return _first_context_number(
            lower,
            [
                r"\b(?:seat(?:ing)?\s*)?(?:capacity|seats?)\s*[:=]?\s*(\d[\d,.]*)",
                r"\b(\d[\d,.]*)\s*capacity\b",
                r"\b(\d[\d,.]*)\s*(?:seats?|spectators|people|guests)\b",
            ],
        )
    if key == "gla":
        return _first_context_number(
            lower,
            [
                r"\b(?:gla|abl|gross\s+leasable\s+area|gross\s+lettable\s+area|gross\s+leaseable\s+area|retail\s+gfa|retail\s+gross\s+floor\s+area|retail\s+floor\s+area|area\s+bruta\s+locavel|área\s+bruta\s+locável)\s*(?:\([^)]*\))?\s*(?:in\s*)?(?:sqm|sq\s*m|m2|m²|square\s+meters?)?\s*[:=]?\s*(\d[\d,.]*)",
                r"\b(\d[\d,.]*)\s*(?:sqm|sq\s*m|m2|m²|square\s+meters?)\s*(?:gla|abl|gross\s+leasable\s+area|gross\s+lettable\s+area|retail\s+gfa|retail\s+gross\s+floor\s+area|area\s+bruta\s+locavel|área\s+bruta\s+locável)\b",
            ],
        )
    if key in {"exhibition_area", "meeting_area"}:
        return _first_context_number(
            lower,
            [
                r"\b(?:exhibition|meeting|event|conference)\s+(?:area|space)\s*[:=]?\s*(\d[\d,.]*)\s*(?:sqm|sq\s*m|m2|m²|square\s+meters?)?",
                r"\b(\d[\d,.]*)\s*(?:sqm|sq\s*m|m2|m²|square\s+meters?)\s*(?:exhibition|meeting|event|conference)\s+(?:area|space)\b",
            ],
        )
    if key in {"annual_passenger_throughput", "passenger_throughput"}:
        return _passenger_throughput_value(lower)
    if key in {"mosque_area", "gross_floor_area", "prayer_hall_area", "site_area", "built_up_area"}:
        return _first_context_number(
            lower,
            [
                r"\b(?:mosque|gross\s+floor|built[- ]?up|building|prayer\s+hall|prayer|site)\s+(?:area|space)\s*[:=]?\s*(\d[\d,.]*)\s*(?:sqm|sq\s*m|m2|m²|square\s+meters?)?",
                r"\b(?:area|floor\s+area)\s*[:=]?\s*(\d[\d,.]*)\s*(?:sqm|sq\s*m|m2|m²|square\s+meters?)?",
                r"\b(\d[\d,.]*)\s*(?:sqm|sq\s*m|m2|m²|square\s+meters?)\s*(?:mosque|gross\s+floor|built[- ]?up|building|prayer\s+hall|prayer|site)?\s*(?:area|space)?\b",
                r"(?:مساحة)\s*[:=]?\s*(\d[\d,.]*)\s*(?:متر|م²|متر مربع)",
            ],
        )
    if key in {"annual_visitors", "annual_visits", "annual_footfall", "footfall", "daily_visitors"}:
        return _first_context_number(
            lower,
            [
                r"\b(?:annual|yearly)?\s*(?:visitors?|visits?|footfall|visitor\s+count)\s*[:=]?\s*(\d[\d,.]*)\s*(?:million|millions|mn|thousand|thousands)?",
                r"\b(\d[\d,.]*)\s*(?:million|millions|mn|thousand|thousands)?\s*(?:annual\s+)?(?:visitors?|visits?|tourists?)\b",
                r"(?:عدد\s+الزوار|زوار|زائر)\s*[:=]?\s*(\d[\d,.]*)",
            ],
        )
    if key == "enrollment":
        return _first_context_number(
            lower,
            [
                r"\b(?:enrollment|enrolment|student\s+count|students?)\s*[:=]?\s*(\d[\d,.]*)\s*(?:million|millions|mn|thousand|thousands|k)?",
                r"\b(\d[\d,.]*)\s*(?:million|millions|mn|thousand|thousands|k)?\s*(?:students?|enrolled)\b",
            ],
        )
    return None


def _passenger_throughput_value(lower: str) -> float | None:
    for pattern in [
        (
            r"\b(?:\d{4}\s+)?international\s+passengers?\s*(\d[\d,.]*)"
            r"\s*(?:plus|\+|and)\s*domestic\s+passengers?\s*(\d[\d,.]*)"
        ),
        (
            r"\b(?:\d{4}\s+)?international\s+passenger\s+numbers?\s*(\d[\d,.]*)"
            r"\s*(?:plus|\+|and)\s*domestic\s+passenger\s+numbers?\s*(\d[\d,.]*)"
        ),
    ]:
        match = re.search(pattern, lower, flags=re.I)
        if not match:
            continue
        values = [_parse_number_token(match.group(index)) for index in (1, 2)]
        if all(value is not None for value in values):
            return sum(value for value in values if value is not None)

    patterns = [
        (
            r"\b(?:annual_)?passenger(?:_throughput|_traffic|_count)?\s*[:=]?\s*"
            r"(\d[\d,.]*)\s*(million|millions|mn|mio|m|milhoes|milhões|millones|"
            r"thousand|thousands|k)?\s*(?:guests?\s*/\s*)?(?:passengers?|pax)?"
        ),
        (
            r"\b(?:annual\s+)?(?:passenger(?:s)?|pax)\s*"
            r"(?:traffic|throughput|volume|count|numbers?)?\s*[:=]?\s*"
            r"(\d[\d,.]*)\s*(million|millions|mn|mio|m|thousand|thousands|k)?"
        ),
        (
            r"\b(\d[\d,.]*)\s*(million|millions|mn|mio|m|thousand|thousands|k)?"
            r"\s*(?:guests?\s*/\s*)?(?:passengers|pax)\b"
        ),
        (
            r"\b(?:served|handled|processed)\s+"
            r"(\d[\d,.]*)\s*(million|millions|mn|mio|m|thousand|thousands|k)?"
            r"\s*(?:people|travellers?|travelers?)\b"
        ),
    ]
    for pattern in patterns:
        for match in re.finditer(pattern, lower, flags=re.I):
            raw = match.group(1)
            suffix = match.group(2) if match.lastindex and match.lastindex >= 2 else ""
            if _passenger_number_looks_like_reporting_year(lower, match.start(1), match.end(1)):
                continue
            context = lower[max(0, match.start() - 32) : match.end() + 32]
            if "capacity" in context and "throughput" not in context and "traffic" not in context:
                continue
            value = _parse_number_token_with_scale(raw, suffix or "")
            if value is None:
                continue
            return _apply_number_suffix(value, suffix or "")
    return None


def _passenger_number_looks_like_reporting_year(lower: str, start: int, end: int) -> bool:
    raw = lower[start:end].strip()
    if "," in raw or "." in raw:
        return False
    value = _parse_number_token(raw)
    if value is None or not float(value).is_integer() or not 1900 <= int(value) <= 2100:
        return False
    after = lower[end : end + 96]
    if re.search(r"\d[\d,.]*", after):
        return True
    before = lower[max(0, start - 32) : start]
    around = f"{before} {after}"
    return bool(
        re.search(
            r"\b(?:report|statistics|data|traffic|movement|movements|numbers?|"
            r"calendar|fiscal|ytd|jan|january|feb|february|mar|march|apr|april|"
            r"may|jun|june|jul|july|aug|august|sep|september|oct|october|"
            r"nov|november|dec|december)\b",
            around,
        )
    )


def _apply_number_suffix(value: float, suffix: str) -> float:
    normalized = suffix.casefold().strip()
    if normalized in {"billion", "billions"}:
        return value * 1_000_000_000
    if normalized in {"million", "millions", "mn", "mio", "m", "milhoes", "milhões", "millones"}:
        return value * 1_000_000
    if normalized in {"thousand", "thousands", "k"}:
        return value * 1_000
    return value


def _first_context_number(
    lower: str,
    patterns: list[str],
    *,
    blocked_contexts: tuple[str, ...] = (),
) -> float | None:
    for pattern in patterns:
        for match in re.finditer(pattern, lower, flags=re.I):
            context = lower[max(0, match.start() - 24) : match.end()]
            if any(blocked in context for blocked in blocked_contexts):
                continue
            suffix = lower[match.end(1) : match.end()]
            value = _parse_number_token_with_scale(match.group(1), suffix)
            if value is not None:
                if re.search(r"\b(billion|billions)\b", suffix):
                    value *= 1_000_000_000
                elif re.search(
                    r"\b(million|millions|mn|mio|milhoes|milhões|millones)\b",
                    suffix,
                ):
                    value *= 1_000_000
                elif re.search(r"\b(thousand|thousands)\b", suffix):
                    value *= 1_000
                return value
    return None


def _metric_number_pattern() -> re.Pattern[str]:
    return re.compile(
        r"(?P<number>\d+(?:[,.]\d{3})*(?:\.\d+)?|\d+(?:[,.]\d+)?)"
        r"\s*(?P<scale>million|millions|billion|billions|thousand|thousands|"
        r"mn|mio|milhoes|milhões|millones|k)?",
        flags=re.I,
    )


def _scaled_number(match: re.Match[str]) -> float | None:
    scale = str(match.group("scale") or "").casefold()
    value = _parse_number_token_with_scale(match.group("number"), scale)
    if value is None:
        return None
    if scale in {"billion", "billions"}:
        return value * 1_000_000_000
    if scale in {"million", "millions", "mn", "mio", "milhoes", "milhões", "millones"}:
        return value * 1_000_000
    if scale in {"thousand", "thousands", "k"}:
        return value * 1_000
    return value


def _parse_number_token(raw: str) -> float | None:
    token = raw.strip()
    if "," in token and "." in token:
        normalized = token.replace(",", "")
    elif "," in token:
        parts = token.split(",")
        normalized = token.replace(",", "") if len(parts[-1]) == 3 else token.replace(",", ".")
    elif "." in token and len(token.split(".")[-1]) == 3:
        normalized = token.replace(".", "")
    else:
        normalized = token
    try:
        return float(normalized)
    except ValueError:
        return None


def _parse_number_token_with_scale(raw: str, scale: str) -> float | None:
    """Parse a number before applying an explicit magnitude suffix.

    A single separator followed by three digits is normally ambiguous and is
    treated as a thousands separator by ``_parse_number_token``. Once a value
    explicitly says ``million``/``mn``/``mio`` (or another magnitude), that
    same spelling represents a decimal magnitude: ``43.712 million`` is
    43.712 million, not 43,712 million.
    """
    normalized_scale = str(scale or "").casefold().strip()
    has_scale = bool(
        re.search(
            r"\b(?:billion|billions|million|millions|mn|mio|m|milhoes|"
            r"milhões|millones|thousand|thousands|k)\b",
            normalized_scale,
        )
    )
    token = raw.strip()
    if has_scale and ((token.count(".") == 1 and "," not in token) or (
        token.count(",") == 1 and "." not in token
    )):
        try:
            return float(token.replace(",", "."))
        except ValueError:
            return None
    return _parse_number_token(token)


def _is_descriptive_metric(metric_key: str) -> bool:
    descriptive_tokens = (
        "role",
        "position",
        "brand",
        "tenant_quality",
        "hospital_grade",
        "building_grade",
        "hotel_class",
        "anchor_brands",
        "urban_catchment",
        "core_facility_intensity",
    )
    return any(token in metric_key for token in descriptive_tokens)


def _has_blocked_objective_text(evidence: EvidenceItem) -> bool:
    rules = load_candidate_quality_rules()
    combined = _normalize_metric_key(f"{evidence.field_value} {evidence.assumption_note or ''}")
    return any(
        _normalize_metric_key(token) in combined
        for token in rules.get("objective_evidence_blocked_tokens", [])
    )


def _evidence_type_rank(value: str) -> int:
    return {"Direct": 0, "Proxy": 1, "Inferred": 2}.get(value, 3)


def _source_tier_rank(value: str) -> int:
    return {"Tier 1": 0, "Tier 2": 1, "Tier 3": 2}.get(value, 3)


def _source_date_rank(value: str | None) -> int:
    if not value:
        return 0
    parts = [int(part) for part in re.findall(r"\d+", value)[:3]]
    if not parts:
        return 0
    year = parts[0]
    month = parts[1] if len(parts) > 1 and 1 <= parts[1] <= 12 else 1
    day = parts[2] if len(parts) > 2 and 1 <= parts[2] <= 31 else 1
    return year * 10_000 + month * 100 + day


def _cross_check_rank(value: str) -> int:
    return {
        "Strong Cross-check": 0,
        "Partial Cross-check": 1,
        "Single Source": 2,
        "Not Checked": 3,
        "Evidence Conflict": 4,
    }.get(value, 5)


def _write_country_summary(worksheet, packets: list[SitePacket], locale: str) -> None:
    countries = sorted({packet.entity.country for packet in packets})
    scene_columns = {
        scene_label("airport_terminal", locale): "airport_terminal",
        scene_label("convention_center", locale): "convention_center",
        scene_label("mall_mixed_use", locale): "mall_mixed_use",
        scene_label("stadium", locale): "stadium",
        scene_label("luxury_hotel_mice", locale): "luxury_hotel_mice",
        scene_label("office_government", locale): "office_government",
        scene_label("hospital", locale): "hospital",
        scene_label("university", locale): "university",
        scene_label("transport_hub", locale): "transport_hub",
        scene_label("mosque", locale): "mosque",
    }
    for country in countries:
        country_packets = [packet for packet in packets if packet.entity.country == country]
        worksheet.append(
            [
                country,
                len(country_packets),
                *[
                    sum(packet.entity.scene_type == scene_type for packet in country_packets)
                    for scene_type in scene_columns.values()
                ],
            ]
        )


def _method_rows(scope_text: str, locale: str) -> list[list[str]]:
    rows = excel_labels(locale).get("method_rows") or []
    return [[str(cell).format(scope_text=scope_text) for cell in row] for row in rows]


def _translate_for_output(
    source_text: str | None,
    text_kind: str,
    locale: str,
    localization_cache: DatabaseLocalizationCache | None,
) -> str:
    precleaned = _clean_locale_text(source_text or "", locale, text_kind)
    if locale == "zh" and precleaned and not re.search(r"[\u4e00-\u9fff]", precleaned):
        return precleaned
    if locale == "en" and precleaned and not re.search(r"[\u4e00-\u9fff]", precleaned):
        return precleaned
    translated = localize_text(
        source_text or "",
        text_kind=text_kind,
        target_locale=locale,
        cache=localization_cache,
        allow_provider=False,
        fallback_text=generic_free_text_fallback(text_kind, locale),
    ).translated_text
    cleaned = _clean_locale_text(translated, locale, text_kind)
    pending = generic_free_text_fallback(text_kind, locale)
    if locale == "zh" and (not cleaned or pending in translated):
        return precleaned or (source_text or "")
    return cleaned


def _clean_locale_text(text: str, locale: str, text_kind: str) -> str:
    if not text:
        return text
    if locale != "en":
        pending = generic_free_text_fallback(text_kind, locale)
        return str(text).replace(pending, "").strip()
    cleaned = str(text)
    replacements = {
        "门户机场": "Gateway airport",
        "会展场馆": "Convention venue",
        "大型体育场": "Large stadium",
        "大型商超": "Large retail mall",
        "奢华酒店": "Luxury hotel",
        "医院": "Hospital",
        "大学": "University",
        "交通枢纽": "Transport hub",
        "写字楼": "Office tower",
        "主指标": "Primary metric",
        "补查": "Check",
        "待核验": "pending verification",
        "部分依据推测": "Some basis is speculative",
        "运营商列表": "operator list",
        "无公开证据": "No public evidence",
        "未知": "Unknown",
    }
    for source, target in replacements.items():
        cleaned = cleaned.replace(source, target)
    cleaned = cleaned.replace("（", " (").replace("）", ")").replace("：", ":")
    cleaned = re.sub(r"\bPrimary metric:\s*", "Primary metric: ", cleaned)
    if re.search(r"[\u4e00-\u9fff]", cleaned):
        return _english_output_fallback(text_kind)
    return re.sub(r"\s+", " ", cleaned).strip()


def _english_output_fallback(text_kind: str) -> str:
    if "reason" in text_kind:
        return (
            "High-value candidate based on scene-level objective evidence; "
            "verify operator indoor coverage status."
        )
    if "next_action" in text_kind or "review" in text_kind:
        return (
            "Verify operator indoor coverage announcements, owner network upgrade notices, "
            "and Google Maps coordinates."
        )
    if "evidence" in text_kind or "inference" in text_kind:
        return "Objective evidence retained in the original evidence column."
    return "Evidence-backed candidate; verify source chain before final action."


def _range_text(values: list[float]) -> str:
    return f"{values[0]}-{values[1]}"
